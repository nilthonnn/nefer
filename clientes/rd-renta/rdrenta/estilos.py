"""Estilos reutilizables del formato de despacho y recepcion."""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import layout

_lado = Side(style="thin", color="FF000000")
BORDE = Border(left=_lado, right=_lado, top=_lado, bottom=_lado)

FILL_CABECERA = PatternFill("solid", fgColor=layout.GRIS_CABECERA)
FILL_RECUPERACION = PatternFill("solid", fgColor=layout.AMARILLO_RECUPERACION)

FUENTE_ETIQUETA = Font(name=layout.FUENTE, size=layout.PT_ETIQUETA, bold=True)
FUENTE_VALOR = Font(name=layout.FUENTE, size=layout.PT_VALOR, bold=False)
FUENTE_TITULO = Font(name=layout.FUENTE, size=layout.PT_TITULO, bold=True)
FUENTE_ROTULO = Font(name=layout.FUENTE, size=layout.PT_ROTULO, bold=True)
FUENTE_META = Font(name=layout.FUENTE, size=layout.PT_VALOR, bold=True)

CENTRO = Alignment(horizontal="center", vertical="center")
CENTRO_AJUSTADO = Alignment(horizontal="center", vertical="center", wrap_text=True)
IZQUIERDA = Alignment(horizontal="left", vertical="center")
IZQUIERDA_AJUSTADA = Alignment(horizontal="left", vertical="center", wrap_text=True)


def pintar_rango(ws, rango: str, *, borde=BORDE, fill=None):
    """Aplica borde (y relleno) a todas las celdas de un rango, incluso combinadas."""
    celdas = ws[rango]
    if not isinstance(celdas, tuple):   # rango de una sola celda
        celdas = ((celdas,),)
    for fila in celdas:
        for celda in fila:
            if borde is not None:
                celda.border = borde
            if fill is not None:
                celda.fill = fill


def escribir(ws, rango: str, valor, *, fuente=FUENTE_VALOR, alineacion=CENTRO,
             fill=None, borde=BORDE, formato=None):
    """Combina `rango`, escribe `valor` en su ancla y aplica el estilo completo."""
    ancla = rango.split(":")[0]
    if ":" in rango:
        ws.merge_cells(rango)
    pintar_rango(ws, rango, borde=borde, fill=fill)
    celda = ws[ancla]
    celda.value = valor
    celda.font = fuente
    celda.alignment = alineacion
    if fill is not None:
        celda.fill = fill
    if formato:
        celda.number_format = formato
    return celda
