"""Ida y vuelta: manifiesto -> Excel -> manifiesto."""

import json

import openpyxl
import pytest

from nefer import build, extract, layout, schema


def foto_falsa(ruta, color=(200, 120, 40)):
    Image = pytest.importorskip("PIL.Image", reason="Pillow es necesario para las imagenes")
    Image.new("RGB", (640, 480), color).save(ruta)


@pytest.fixture
def manifiesto(tmp_path):
    fotos = tmp_path / "fotos"
    fotos.mkdir()
    vistas = ["VISTA FRONTAL", "VISTA POSTERIOR", "VISTA LATERAL IZQUIERDA",
              "VISTA LATERAL DERECHA", "HORÓMETRO", "PANEL DE CONTROL"]
    registro = []
    for i, vista in enumerate(vistas, start=1):
        nombre = f"foto{i}.png"
        foto_falsa(fotos / nombre)
        registro.append({"foto_id": i, "descripcion": vista, "archivo": f"fotos/{nombre}"})

    foto_falsa(fotos / "extintor.png", (200, 30, 30))
    return {
        "encabezado": {
            "empresa": "Maquinarias del Sur S.A.C.",
            "tipo_documento": "RECEPCION",
            "n_acta": "031-004000",
            "n_guia": "EG07-00009999",
            "cliente": "CONSTRUCTORA DEMO S.A.C.",
            "obra": "OBRA DEMO",
            "fecha": "2026-03-15",
            "horometro": 2450.5,
            "codigo_equipo": "GE074-99",
            "modelo_equipo": "GRUPO ELECTRÓGENO INSONORIZADO DE 74 KW",
            "categoria": "grupo_electrogeno",
        },
        "inspeccion_componentes": [
            {"item": "Módulo de control", "estado": "OK", "observacion": ""},
            {"item": "Nivel de combustible", "estado": "OBS",
             "observacion": "Tanque al 45%."},
        ],
        "registro_fotografico": registro,
        "consumibles": [{
            "descripcion": "EXTINTOR DE 6 KG",
            "cantidad": 1,
            "estado_recepcion": "NO_RETORNA",
            "foto_despacho": "fotos/extintor.png",
        }],
        "control_consumibles": [
            {"consumible": "Combustible diésel", "unidad": "%",
             "despacho": 100, "recepcion": 45},
        ],
        "resumen_ejecutivo": "Equipo operativo; retorna sin extintor, se genera recuperacion.",
    }


def test_manifiesto_de_prueba_es_valido(manifiesto, tmp_path):
    assert schema.validar(manifiesto, tmp_path) == []


def test_construye_las_hojas_esperadas(manifiesto, tmp_path):
    salida, avisos = build.construir(manifiesto, tmp_path / "acta.xlsx", raiz=tmp_path)
    assert avisos == []
    wb = openpyxl.load_workbook(salida)
    assert wb.sheetnames == ["REPORTE", "INSPECCIÓN", "CONSUMIBLES",
                             "GUÍA OPERADOR", "GUÍA CLIENTE"]


def test_cabecera_en_las_celdas_del_formato(manifiesto, tmp_path):
    salida, _ = build.construir(manifiesto, tmp_path / "acta.xlsx", raiz=tmp_path)
    ws = openpyxl.load_workbook(salida)["REPORTE"]
    assert ws[layout.CELDA_ACTA].value == "031-004000"
    assert ws[layout.CELDA_CODIGO].value == "GE074-99"
    assert ws[layout.CELDA_HOROMETRO].value == 2450.5
    # Es una recepcion: la "X" va en la casilla derecha, no en la de despacho.
    assert ws[layout.CELDA_MARCA_RECEPCION].value == "X"
    assert ws[layout.CELDA_MARCA_DESPACHO].value is None


def test_las_fotos_quedan_ancladas_en_su_bloque(manifiesto, tmp_path):
    salida, _ = build.construir(manifiesto, tmp_path / "acta.xlsx", raiz=tmp_path)
    ws = openpyxl.load_workbook(salida)["REPORTE"]
    # 6 fotos + 1 consumible. Sin logo: el manifiesto no declara ninguno.
    assert len(ws._images) == 7
    filas = {img.anchor._from.row + 1 for img in ws._images}
    assert layout.bloque_foto(0)["fila_imagen_inicio"] in filas
    assert layout.bloque_foto(2)["fila_imagen_inicio"] in filas


def test_el_logo_es_opcional_y_configurable(manifiesto, tmp_path):
    """El paquete no trae logo: lo aporta quien usa la herramienta."""
    salida, _ = build.construir(manifiesto, tmp_path / "sin-logo.xlsx", raiz=tmp_path)
    assert len(openpyxl.load_workbook(salida)["REPORTE"]._images) == 7

    foto_falsa(tmp_path / "fotos" / "logo.png", (10, 60, 140))
    manifiesto["encabezado"]["logo"] = "fotos/logo.png"
    salida, _ = build.construir(manifiesto, tmp_path / "con-logo.xlsx", raiz=tmp_path)
    imagenes = openpyxl.load_workbook(salida)["REPORTE"]._images
    assert len(imagenes) == 8
    assert any(img.anchor._from.row == 0 for img in imagenes), "el logo va en la fila 1"


def test_control_documental_configurable(manifiesto, tmp_path):
    """Cada organizacion pone su propio codigo, version y fecha de formato."""
    manifiesto["encabezado"].update(
        {"codigo_formato": "PROC-045", "version_formato": "03",
         "fecha_formato": "01/02/2026"})
    salida, _ = build.construir(manifiesto, tmp_path / "acta.xlsx", raiz=tmp_path)
    ws = openpyxl.load_workbook(salida)["REPORTE"]
    assert ws["S1"].value == "CÓDIGO: PROC-045"
    assert ws["S2"].value == "VERSIÓN: 03"
    assert ws["S3"].value == "FECHA: 01/02/2026"


def test_consumo_calculado_en_la_hoja_de_consumibles(manifiesto, tmp_path):
    salida, _ = build.construir(manifiesto, tmp_path / "acta.xlsx", raiz=tmp_path)
    ws = openpyxl.load_workbook(salida)["CONSUMIBLES"]
    fila = [c.value for c in ws[5]]          # fila 4 = encabezados, 5 = primer dato
    assert fila[3] == 100 and fila[4] == 45 and fila[5] == 55


def test_ida_y_vuelta_conserva_los_datos(manifiesto, tmp_path):
    salida, _ = build.construir(manifiesto, tmp_path / "acta.xlsx", raiz=tmp_path)
    recuperado = extract.extraer(salida)

    original = manifiesto["encabezado"]
    vuelta = recuperado["encabezado"]
    for campo in ("tipo_documento", "n_acta", "n_guia", "fecha",
                  "horometro", "codigo_equipo", "categoria"):
        assert vuelta[campo] == original[campo], campo
    assert vuelta["cliente"].strip() == original["cliente"]

    assert [f["descripcion"] for f in recuperado["registro_fotografico"]] == \
           [f["descripcion"] for f in manifiesto["registro_fotografico"]]
    assert len(recuperado["consumibles"]) == 1
    assert recuperado["consumibles"][0]["estado_recepcion"] == "NO_RETORNA"


def test_horometro_ilegible_se_imprime_como_leyenda(manifiesto, tmp_path):
    manifiesto["encabezado"]["horometro"] = schema.REVISION_MANUAL
    salida, _ = build.construir(manifiesto, tmp_path / "acta.xlsx", raiz=tmp_path)
    ws = openpyxl.load_workbook(salida)["REPORTE"]
    assert ws[layout.CELDA_HOROMETRO].value == schema.REVISION_MANUAL
    assert extract.extraer(salida)["encabezado"]["horometro"] == schema.REVISION_MANUAL


def test_imagen_no_soportada_avisa_sin_romper(manifiesto, tmp_path):
    (tmp_path / "fotos" / "roto.tiff").write_bytes(b"no soy una imagen")
    manifiesto["registro_fotografico"][0]["archivo"] = "fotos/roto.tiff"
    _, avisos = build.construir(manifiesto, tmp_path / "acta.xlsx", raiz=tmp_path)
    assert len(avisos) == 1 and "roto.tiff" in avisos[0]


def test_cli_construir_y_validar(manifiesto, tmp_path):
    from nefer.cli import main

    ruta = tmp_path / "acta.json"
    ruta.write_text(json.dumps(manifiesto, ensure_ascii=False), encoding="utf-8")
    assert main(["validar", str(ruta)]) == 0
    assert main(["construir", str(ruta), "-o", str(tmp_path / "salida.xlsx")]) == 0
    assert (tmp_path / "salida.xlsx").exists()


def test_validar_comprueba_las_fotos_por_defecto(manifiesto, tmp_path, capsys):
    """Decir «válido» y que luego `construir` falle es peor que no validar."""
    from nefer.cli import main

    manifiesto["registro_fotografico"][0]["archivo"] = "fotos/no-esta.jpg"
    ruta = tmp_path / "acta.json"
    ruta.write_text(json.dumps(manifiesto, ensure_ascii=False), encoding="utf-8")

    assert main(["validar", str(ruta)]) == 1
    assert "no-esta.jpg" in capsys.readouterr().err
    # El manifiesto sigue siendo correcto si no se mira el disco.
    assert main(["validar", str(ruta), "--sin-verificar-fotos"]) == 0


def test_json_corrupto_da_un_mensaje_no_una_traza(tmp_path, capsys):
    from nefer.cli import main

    ruta = tmp_path / "roto.json"
    ruta.write_text("{ roto", encoding="utf-8")
    assert main(["validar", str(ruta)]) == 1
    salida = capsys.readouterr().err
    assert "no es JSON valido" in salida and "linea 1" in salida
    assert "Traceback" not in salida


def test_manifiesto_inexistente_da_un_mensaje(tmp_path, capsys):
    from nefer.cli import main

    assert main(["validar", str(tmp_path / "no-existe.json")]) == 1
    assert "No existe el manifiesto" in capsys.readouterr().err


def test_logo_declarado_pero_inexistente_avisa(manifiesto, tmp_path):
    """Un logo que no está no puede desaparecer en silencio."""
    manifiesto["encabezado"]["logo"] = "fotos/no-existe.png"
    _, avisos = build.construir(manifiesto, tmp_path / "acta.xlsx", raiz=tmp_path)
    assert len(avisos) == 1 and "no existe" in avisos[0]


def test_acta_de_despacho_deja_la_recepcion_en_blanco(manifiesto, tmp_path):
    manifiesto["encabezado"]["tipo_documento"] = "DESPACHO"
    manifiesto["consumibles"] = [{"descripcion": "EXTINTOR DE 6 KG", "cantidad": 1}]
    assert schema.validar(manifiesto, tmp_path) == []

    salida, _ = build.construir(manifiesto, tmp_path / "despacho.xlsx", raiz=tmp_path)
    ws = openpyxl.load_workbook(salida)["REPORTE"]
    bloque = layout.bloque_consumible(0, len(manifiesto["registro_fotografico"]) // 2)

    izquierda = ws[f"{layout.PANEL_IZQ[0]}{bloque['fila_rotulo']}"].value
    derecha = ws[f"{layout.PANEL_DER[0]}{bloque['fila_rotulo']}"].value
    franja = ws[f"A{bloque['fila_recuperacion']}"].value
    assert izquierda == "01 EXTINTOR DE 6 KG DESPACHADO"
    assert not derecha, "el lado de recepción va en blanco hasta que el equipo vuelva"
    assert not franja, "no hay recuperación que declarar en un despacho"
    assert ws[layout.CELDA_MARCA_DESPACHO].value == "X"


def test_ida_y_vuelta_recupera_las_fotos(manifiesto, tmp_path):
    """El extractor debe leer las actas que genera este mismo paquete.

    openpyxl declara el namespace de dibujo por defecto y escribe `<from>`,
    mientras que Excel escribe `<xdr:from>`. Leer solo la variante con prefijo
    dejaba el round-trip fotográfico en cero sin fallar ni avisar.
    """
    manifiesto["encabezado"]["logo"] = manifiesto["registro_fotografico"][0]["archivo"]
    salida, _ = build.construir(manifiesto, tmp_path / "acta.xlsx", raiz=tmp_path)

    destino = tmp_path / "recuperadas"
    recuperado = extract.extraer(salida, destino)

    fotos = recuperado["registro_fotografico"]
    assert len(fotos) == len(manifiesto["registro_fotografico"])
    assert all(f.get("archivo") for f in fotos), "toda foto debe volver con su archivo"
    assert recuperado["consumibles"][0].get("foto_despacho")

    # 6 fotos + 1 consumible. El logo va en la fila 1 y no es una vista del equipo.
    assert len(list(destino.glob("*"))) == 7


def test_las_anclas_se_leen_con_y_sin_prefijo_xdr(manifiesto, tmp_path):
    import zipfile

    salida, _ = build.construir(manifiesto, tmp_path / "acta.xlsx", raiz=tmp_path)
    with zipfile.ZipFile(salida) as z:
        anclas = extract._anclas_de_imagen(z)
    assert len(anclas) == 7
    assert all(isinstance(c, int) and isinstance(f, int) for c, f, _ in anclas)


def _carpeta_fotos(tmp_path, nombres):
    carpeta = tmp_path / "fotos"
    carpeta.mkdir(exist_ok=True)
    for n in nombres:
        foto_falsa(carpeta / n)
    return carpeta


def test_fotos_empareja_por_orden_de_nombre(tmp_path, capsys):
    """El orden alfabético del archivo es el orden de la rejilla."""
    from nefer.cli import main

    _carpeta_fotos(tmp_path, ["03-c.png", "01-a.png", "02-b.png"])
    assert main(["fotos", str(tmp_path / "fotos"), "-c", "grupo_electrogeno"]) == 0

    bloque = json.loads(capsys.readouterr().out)["registro_fotografico"]
    assert [f["archivo"].split("/")[-1] for f in bloque] == \
           ["01-a.png", "02-b.png", "03-c.png"]
    assert [f["descripcion"] for f in bloque][:2] == ["VISTA FRONTAL", "VISTA POSTERIOR"]
    assert [f["foto_id"] for f in bloque] == [1, 2, 3]


def test_fotos_escribe_en_el_manifiesto_y_usa_su_categoria(manifiesto, tmp_path):
    from nefer.cli import main

    manifiesto["encabezado"]["categoria"] = "plataforma_elevacion"
    ruta = tmp_path / "acta.json"
    ruta.write_text(json.dumps(manifiesto, ensure_ascii=False), encoding="utf-8")
    _carpeta_fotos(tmp_path, ["a.png", "b.png"])

    assert main(["fotos", str(tmp_path / "fotos"), "-m", str(ruta)]) == 0
    vuelta = json.loads(ruta.read_text(encoding="utf-8"))
    bloque = vuelta["registro_fotografico"]
    # Rótulos de plataforma, y rutas relativas al manifiesto.
    assert bloque[0]["descripcion"] == "VISTA FRONTAL"
    assert bloque[0]["archivo"] == "fotos/a.png"
    assert schema.validar(vuelta, tmp_path) == []


def test_fotos_avisa_de_los_formatos_que_excel_no_incrusta(tmp_path, capsys):
    """Un HEIC de iPhone se detecta antes de generar, no después."""
    from nefer.cli import main

    carpeta = _carpeta_fotos(tmp_path, ["a.png"])
    (carpeta / "IMG_0042.HEIC").write_bytes(b"no importa")

    assert main(["fotos", str(carpeta), "-c", "generico"]) == 0
    err = capsys.readouterr().err
    assert "IMG_0042.HEIC" in err and "Mas compatible" in err


def test_fotos_avisa_cuando_sobran_fotos(tmp_path, capsys):
    from nefer.cli import main

    _carpeta_fotos(tmp_path, [f"{i:02d}.png" for i in range(8)])
    assert main(["fotos", str(tmp_path / "fotos"), "-c", "generico"]) == 0
    assert "solo 6 rotulos" in capsys.readouterr().err


def test_fotos_carpeta_inexistente(tmp_path, capsys):
    from nefer.cli import main

    assert main(["fotos", str(tmp_path / "no-esta")]) == 1
    assert "No existe la carpeta" in capsys.readouterr().err


def test_fotos_usa_las_pistas_del_nombre(tmp_path, capsys):
    """Un nombre descriptivo manda sobre el orden alfabético."""
    from nefer.cli import main

    # 'panel' va antes que 'zz-frontal' alfabéticamente, pero cada una
    # debe caer en la casilla que su nombre indica.
    _carpeta_fotos(tmp_path, ["panel.png", "zz-frontal.png"])
    assert main(["fotos", str(tmp_path / "fotos"), "-c", "grupo_electrogeno"]) == 0

    bloque = json.loads(capsys.readouterr().out)["registro_fotografico"]
    por_vista = {f["descripcion"]: f["archivo"].split("/")[-1] for f in bloque}
    assert por_vista["VISTA FRONTAL"] == "zz-frontal.png"
    assert por_vista["PANEL DE CONTROL"] == "panel.png"


def test_fotos_rellena_los_huecos_con_las_opacas(tmp_path, capsys):
    """Las que no dicen nada ocupan las casillas que quedaron libres."""
    from nefer.cli import main

    _carpeta_fotos(tmp_path, ["IMG_001.png", "IMG_002.png", "baterias.png"])
    assert main(["fotos", str(tmp_path / "fotos"), "-c", "grupo_electrogeno"]) == 0

    bloque = json.loads(capsys.readouterr().out)["registro_fotografico"]
    por_vista = {f["descripcion"]: f["archivo"].split("/")[-1] for f in bloque}
    assert por_vista["BATERÍAS"] == "baterias.png"
    assert por_vista["VISTA FRONTAL"] == "IMG_001.png"
    assert por_vista["VISTA POSTERIOR"] == "IMG_002.png"


def test_vista_sugerida_ignora_tildes_y_mayusculas():
    assert layout.vista_sugerida("05-HORÓMETRO.JPG") == "HORÓMETRO"
    assert layout.vista_sugerida("Lat_Izq.jpeg") == "VISTA LATERAL IZQUIERDA"
    assert layout.vista_sugerida("IMG_4471.JPG") is None


# --------------------------------------------------------------------------- #
# Catalogo de clientes y equipos
# --------------------------------------------------------------------------- #
def _catalogo(tmp_path):
    from nefer import catalogo

    datos = {
        "empresa": "Maquinarias del Sur S.A.C.",
        "codigo_formato": "MDS-FO-001", "version_formato": "02",
        "clientes": [
            {"id": "andina", "razon_social": "CONSTRUCTORA ANDINA S.A.C.",
             "obras": ["PLANTA CONCENTRADORA — FASE II"]},
            {"id": "pacifico", "razon_social": "MONTAJES DEL PACÍFICO E.I.R.L.",
             "obras": ["OBRA A", "OBRA B"]},
        ],
        "equipos": [
            {"codigo": "GE110-02", "modelo": "GRUPO ELECTRÓGENO DE 110 KW",
             "categoria": "grupo_electrogeno"},
            {"codigo": "TI009-04", "modelo": "TORRE DE ILUMINACIÓN 4x1000 W",
             "categoria": "torre_iluminacion"},
        ],
    }
    ruta = tmp_path / "catalogo.json"
    ruta.write_text(json.dumps(datos, ensure_ascii=False), encoding="utf-8")
    return catalogo.cargar(ruta)


def test_catalogo_llena_el_encabezado_entero(tmp_path):
    from nefer import catalogo

    enc = catalogo.encabezado(_catalogo(tmp_path), "GE110-02", "andina", "DESPACHO")
    assert enc["empresa"] == "Maquinarias del Sur S.A.C."
    assert enc["cliente"] == "CONSTRUCTORA ANDINA S.A.C."
    assert enc["obra"] == "PLANTA CONCENTRADORA — FASE II"
    assert enc["modelo_equipo"] == "GRUPO ELECTRÓGENO DE 110 KW"
    assert enc["categoria"] == "grupo_electrogeno"
    assert enc["codigo_formato"] == "MDS-FO-001"


def test_catalogo_busca_por_fragmento_y_sin_tildes(tmp_path):
    from nefer import catalogo

    cat = _catalogo(tmp_path)
    assert catalogo.equipo(cat, "ge110")["codigo"] == "GE110-02"
    assert catalogo.equipo(cat, "torre")["codigo"] == "TI009-04"
    assert catalogo.cliente(cat, "PACIFICO")["id"] == "pacifico"


def test_catalogo_no_adivina_cuando_hay_varias_coincidencias(tmp_path):
    from nefer import catalogo

    cat = _catalogo(tmp_path)
    cat["equipos"].append({"codigo": "GE110-99", "modelo": "OTRO",
                           "categoria": "grupo_electrogeno"})
    with pytest.raises(catalogo.ErrorCatalogo, match="varios"):
        catalogo.equipo(cat, "GE110")


def test_catalogo_dice_que_hay_disponible_si_no_encuentra(tmp_path):
    from nefer import catalogo

    with pytest.raises(catalogo.ErrorCatalogo, match="Disponibles"):
        catalogo.equipo(_catalogo(tmp_path), "NO-EXISTE")


def test_obra_no_se_adivina_si_el_cliente_tiene_varias(tmp_path):
    from nefer import catalogo

    enc = catalogo.encabezado(_catalogo(tmp_path), "GE110-02", "pacifico")
    assert enc["obra"] == "", "con dos obras posibles no se elige ninguna"


def test_cmd_acta_deja_un_manifiesto_casi_completo(tmp_path, monkeypatch, capsys):
    from nefer.cli import main

    _catalogo(tmp_path)
    monkeypatch.chdir(tmp_path)
    assert main(["acta", "-e", "ge110", "-c", "andina", "--acta", "004-001"]) == 0

    m = json.loads((tmp_path / "acta.json").read_text(encoding="utf-8"))
    enc = m["encabezado"]
    assert enc["codigo_equipo"] == "GE110-02"
    assert enc["n_acta"] == "004-001"
    assert enc["fecha"]                      # hoy, por defecto
    # Lo que el catálogo no puede saber queda pendiente y el aviso lo dice.
    assert enc["horometro"] == 0.0
    assert "REVISIÓN MANUAL REQUERIDA" in capsys.readouterr().out


def test_cmd_acta_no_pisa_un_acta_existente(tmp_path, monkeypatch, capsys):
    from nefer.cli import main

    _catalogo(tmp_path)
    monkeypatch.chdir(tmp_path)
    (tmp_path / "acta.json").write_text("{}", encoding="utf-8")
    assert main(["acta", "-e", "ge110"]) == 1
    assert "Ya existe" in capsys.readouterr().err
    assert main(["acta", "-e", "ge110", "--forzar"]) == 0


def test_fecha_sale_del_exif_de_las_fotos(manifiesto, tmp_path):
    """La fecha del acta es la de captura, no la de generación."""
    from PIL import Image
    from nefer.cli import main

    for nombre in ("a.png", "b.png"):
        foto_falsa(tmp_path / "fotos" / nombre)
    for nombre, cuando in (("c.jpg", "2026:07:14 07:30:00"),
                           ("d.jpg", "2026:07:14 07:31:00")):
        exif = Image.Exif()
        exif[0x0132] = cuando
        Image.new("RGB", (64, 48), (10, 20, 30)).save(tmp_path / "fotos" / nombre, exif=exif)

    manifiesto["encabezado"]["fecha"] = ""
    ruta = tmp_path / "acta.json"
    ruta.write_text(json.dumps(manifiesto, ensure_ascii=False), encoding="utf-8")

    assert main(["fotos", str(tmp_path / "fotos"), "-m", str(ruta)]) == 0
    vuelta = json.loads(ruta.read_text(encoding="utf-8"))
    assert vuelta["encabezado"]["fecha"] == "2026-07-14"


# --------------------------------------------------------------------------- #
# Secciones pareadas del acta de recepción
# --------------------------------------------------------------------------- #
def _recepcion_con_pares(manifiesto, tmp_path):
    """Acta de recepción con fotos de antes y después, y un componente dañado."""
    fotos = tmp_path / "fotos"
    for n in ("antes-frontal.png", "antes-panel.png", "antes-escape.png",
              "despues-escape.png"):
        foto_falsa(fotos / n)

    manifiesto["encabezado"]["tipo_documento"] = "RECEPCION"
    manifiesto["registro_fotografico"][0]["archivo_despacho"] = "fotos/antes-frontal.png"
    manifiesto["registro_fotografico"][5]["archivo_despacho"] = "fotos/antes-panel.png"
    manifiesto["inspeccion_componentes"] = [
        {"item": "Módulo de control", "estado": "OK", "observacion": ""},
        {"item": "Sistema de escape", "estado": "D",
         "observacion": "Junta rota; se factura el cambio.",
         "foto_despacho": "fotos/antes-escape.png",
         "foto_recepcion": "fotos/despues-escape.png"},
    ]
    return manifiesto


def test_con_fotos_de_salida_el_informe_va_emparejado(manifiesto, tmp_path):
    """El informe fotográfico de la recepción es la tabla del formato."""
    from nefer import build

    m = _recepcion_con_pares(manifiesto, tmp_path)
    assert schema.validar(m, tmp_path) == []
    entradas = build.entradas_vistas(m)
    assert len(entradas) == len(m["registro_fotografico"]), \
        "todas las vistas, no sólo las que traen foto de salida"
    # El rótulo va igual en las dos columnas: quien manda es el encabezado.
    assert entradas[0]["texto_izq"] == entradas[0]["texto_der"] == \
        m["registro_fotografico"][0]["descripcion"]
    assert entradas[0]["foto_izq"] == "fotos/antes-frontal.png"
    assert entradas[0]["foto_der"] == m["registro_fotografico"][0]["archivo"]
    # La vista sin foto de salida entra igual, con su celda izquierda vacía.
    assert entradas[1]["foto_izq"] is None


def test_sin_fotos_de_salida_la_recepcion_conserva_la_rejilla(manifiesto, tmp_path):
    """Sin nada que emparejar, la rejilla del formato ocupa la mitad de hojas."""
    from nefer import build

    m = _recepcion_con_pares(manifiesto, tmp_path)
    for vista in m["registro_fotografico"]:
        vista.pop("archivo_despacho", None)
    assert build.entradas_vistas(m) == []


def test_danos_solo_con_lo_observado_o_danado(manifiesto, tmp_path):
    from nefer import build

    entradas = build.entradas_danos(_recepcion_con_pares(manifiesto, tmp_path))
    assert len(entradas) == 1, "el componente OK no entra"
    assert "Sistema de escape".upper() in entradas[0]["rotulo"]
    assert "DAÑADO" in entradas[0]["texto_der"]
    assert entradas[0]["pie"].startswith("Junta rota")


def test_un_despacho_no_lleva_secciones_pareadas(manifiesto, tmp_path):
    from nefer import build

    m = _recepcion_con_pares(manifiesto, tmp_path)
    m["encabezado"]["tipo_documento"] = "DESPACHO"
    assert build.entradas_vistas(m) == []
    assert build.entradas_danos(m) == []


def test_despacho_rechaza_los_campos_de_comparacion(manifiesto, tmp_path):
    m = _recepcion_con_pares(manifiesto, tmp_path)
    m["encabezado"]["tipo_documento"] = "DESPACHO"
    errores = schema.validar(m, tmp_path)
    assert any("archivo_despacho: solo tiene sentido" in e for e in errores)
    assert any("foto_recepcion: un acta de DESPACHO" in e for e in errores)


def test_las_observaciones_van_pegadas_al_informe_fotografico(manifiesto, tmp_path):
    """Sin sección intermedia: las fotos del acta se leen seguidas.

    Antes había una sección COMPARATIVO entre las observaciones y las fotos con
    las que se leen, y cada foto del retorno se imprimía dos veces —una en la
    rejilla y otra en el comparativo—.
    """
    from nefer import build

    m = _recepcion_con_pares(manifiesto, tmp_path)
    n_vistas = len(m["registro_fotografico"])
    salida, avisos = build.construir(m, tmp_path / "recepcion.xlsx", raiz=tmp_path)
    assert avisos == []
    ws = openpyxl.load_workbook(salida)["REPORTE"]

    columna_a = [ws.cell(row=f, column=1).value for f in range(1, ws.max_row + 1)]
    assert layout.TITULO_COMPARATIVO not in columna_a, "esa sección ya no existe"

    titulos = {v: f + 1 for f, v in enumerate(columna_a)
               if v in (layout.TITULO_DANOS, "OBSERVACIONES")}
    assert titulos["OBSERVACIONES"] < titulos[layout.TITULO_DANOS]

    # El informe emparejado abre el acta, y OBSERVACIONES viene justo después.
    ultima_vista = layout.bloque_pareado(n_vistas - 1, layout.FILA_VISTAS)
    assert ws[f"A{layout.FILA_INICIO_FOTOS}"].value == "DESPACHO"
    assert titulos["OBSERVACIONES"] == ultima_vista["fila_pie"] + 1

    # 6 vistas con su foto de retorno + 2 fotos de salida + 1 consumible +
    # 1 par de daños. Ninguna repetida.
    assert len(ws._images) == 6 + 2 + 1 + 2
    nombres = [im.ref.name if hasattr(im.ref, "name") else str(im.ref)
               for im in ws._images]
    assert len(nombres) == len(set(nombres)), "ninguna foto se imprime dos veces"


def test_la_extraccion_no_confunde_los_pares_con_consumibles(manifiesto, tmp_path):
    """Las secciones pareadas usan los mismos rótulos DESPACHO / RECEPCIÓN."""
    from nefer import build

    m = _recepcion_con_pares(manifiesto, tmp_path)
    salida, _ = build.construir(m, tmp_path / "recepcion.xlsx", raiz=tmp_path)
    recuperado = extract.extraer(salida)
    assert len(recuperado["consumibles"]) == 1, \
        "los bloques comparativos no son consumibles"
    assert recuperado["consumibles"][0]["descripcion"] == "EXTINTOR DE 6 KG"


def test_el_retorno_parcial_ocupa_una_fila_mas_en_el_acta(manifiesto, tmp_path):
    """El bloque de un accesorio que vuelve en parte crece una fila."""
    man = manifiesto
    man["consumibles"] = [
        {"descripcion": "GANCHOS DE IZAJE", "cantidad": 2,
         "estado_recepcion": "NO_RETORNA", "cantidad_retorna": 1},
        {"descripcion": 'CONOS DE 28"', "cantidad": 2,
         "estado_recepcion": "NO_RETORNA"},
    ]
    assert schema.validar(man, tmp_path) == []

    salida = tmp_path / "acta.xlsx"
    build.construir(man, salida, tmp_path)
    ws = openpyxl.load_workbook(salida).active

    bandas = {f: ws.cell(row=f, column=1).value
              for f in range(1, ws.max_row + 1)
              if isinstance(ws.cell(row=f, column=1).value, str)
              and ws.cell(row=f, column=1).value.startswith(("RECUPERACIÓN", "CONFORME"))}
    assert list(bandas.values()) == [
        "RECUPERACIÓN N° 1 : 01 GANCHOS DE IZAJE",
        "CONFORME N° 1 : 01 GANCHOS DE IZAJE — SIN RECUPERACIÓN",
        'RECUPERACIÓN N° 2 : 02 CONOS DE 28"',
    ]
    # Las dos primeras van seguidas, y el bloque siguiente empieza despues.
    filas = sorted(bandas)
    assert filas[1] == filas[0] + 1
    assert filas[2] == filas[1] + layout.ALTO_BLOQUE_CONSUMIBLE


def test_una_cantidad_de_retorno_imposible_se_rechaza(manifiesto, tmp_path):
    man = manifiesto
    man["consumibles"] = [{"descripcion": "GANCHOS", "cantidad": 2,
                           "estado_recepcion": "NO_RETORNA", "cantidad_retorna": 3}]
    errores = schema.validar(man, tmp_path)
    assert any("cantidad_retorna" in e for e in errores), errores


def test_un_despacho_no_puede_declarar_cuantas_volvieron(manifiesto, tmp_path):
    man = manifiesto
    man["encabezado"]["tipo_documento"] = "DESPACHO"
    man["inspeccion_componentes"] = []
    man["consumibles"] = [{"descripcion": "GANCHOS", "cantidad": 2,
                           "cantidad_retorna": 1}]
    errores = schema.validar(man, tmp_path)
    assert any("cantidad_retorna" in e for e in errores), errores
