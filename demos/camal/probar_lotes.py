"""Recorre la app de punta a punta, como una jornada, y audita lo que exporta.

No basta con que la pantalla diga lo correcto: el Excel y el PDF se abren
después con librerías ajenas a la app —openpyxl y pypdf— y se comprueba que
digan lo mismo. Un registro trazable que solo cuadra en su propia pantalla no
sirve de nada.

    python3 demos/camal/probar_lotes.py salidas/
"""

import pathlib
import sys

from playwright.sync_api import sync_playwright

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from _navegador import ruta_del_navegador

APP = pathlib.Path(__file__).resolve().parent / "app-camal.html"
SALIDA = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "salidas")
SALIDA.mkdir(parents=True, exist_ok=True)

errores, fallos = [], []


def exigir(condicion, dicho):
    print(("  ok  " if condicion else "FALLA ") + dicho)
    if not condicion:
        fallos.append(dicho)


with sync_playwright() as pw:
    nav = pw.chromium.launch(executable_path=ruta_del_navegador())
    ctx = nav.new_context(viewport={"width": 390, "height": 844},
                          device_scale_factor=2, locale="es-PE",
                          timezone_id="America/Lima", accept_downloads=True)
    p = ctx.new_page()
    p.on("pageerror", lambda e: errores.append(str(e)))
    p.goto(APP.as_uri())

    print("\n[1] La jornada abre con datos de ejemplo y su control documental")
    exigir("Jornada del camal" in p.inner_text("h1"), "se ve la jornada")
    exigir(p.locator(".lote").count() == 3, "tres lotes de demostración")
    exigir(p.evaluate("jornada.id").startswith("J-"), "la jornada tiene identificador")
    exigir("REG-CAM-001" in p.inner_text(".ficha-control"), "el formato está a la vista")
    p.fill("#operador", "Nilthon Chit")
    p.dispatch_event("#operador", "change")
    exigir(p.evaluate("jornada.operador") == "Nilthon Chit", "queda el responsable del registro")
    p.screenshot(path=SALIDA / "1-jornada.png", full_page=True)

    print("\n[2] Lote nuevo de pesaje: 40 alpacas degolladas")
    p.click("text=Nuevo lote de pesaje")
    p.select_option("#categoriaLote", "Alpaca degollada")
    exigir(p.input_value("#cantidadLote") == "40", "la cantidad propuesta es 40")
    p.click("text=Empezar a pesar")
    exigir("de 40" in p.inner_text(".correlativo"), "arranca en el correlativo 1 de 40")
    exigir(p.evaluate("codigoLote(loteActual())").endswith("-L004"),
           "el lote lleva código correlativo de tres cifras")
    p.screenshot(path=SALIDA / "2-pesaje-inicio.png", full_page=True)

    print("\n[3] Se pesa y salta solo al siguiente correlativo")
    for peso in ["42.5", "39,8", "44.2"]:      # el segundo va con coma, a propósito
        p.fill("#pesoActual", peso)
        p.click("text=/^Registrar N°/")
    exigir(p.inner_text(".correlativo .num") == "4", "tras tres pesadas toca el N° 4")
    exigir(p.evaluate("loteActual().pesos[1].kg") == 39.8, "«39,8» con coma entra como 39.8")
    exigir("126.50" in p.inner_text(".tiras"), "acumula 126.50 kg")
    p.screenshot(path=SALIDA / "3-pesaje-tres.png", full_page=True)

    print("\n[4] Lo que no es un peso no entra")
    for malo in ["asd", "-12", "0", "980"]:
        p.fill("#pesoActual", malo)
        p.click("text=/^Registrar N°/")
        exigir(p.locator(".aviso").count() == 1 and p.evaluate("loteActual().pesos.length") == 3,
               f"«{malo}» se rechaza con aviso")

    print("\n[5] La cantidad del lote se amplía y no baja de lo pesado")
    p.click("button[aria-label='Sumar una cabeza']")
    exigir(p.evaluate("loteActual().cantidad") == 41, "sube a 41")
    p.fill("#cantidadEnCurso", "2")
    p.dispatch_event("#cantidadEnCurso", "change")
    exigir(p.evaluate("loteActual().cantidad") == 3, "no baja de las 3 pesadas")

    print("\n[6] Corregir y anular exigen motivo, y no borran nada")
    p.click(".pesada >> nth=0")
    p.fill("#pesoEditado", "41,0")
    p.click(".gaveta button.btn:not(.btn-rojo):not(.btn-sec)")
    exigir(p.evaluate("loteActual().pesos[0].kg") == 42.5, "sin motivo, la corrección no se aplica")
    exigir(p.locator(".aviso").count() == 1, "y lo dice con un aviso")

    p.click(".pesada >> nth=0")
    p.fill("#pesoEditado", "41,0")
    p.fill("#motivoCambio", "balanza mal tarada")
    p.click(".gaveta button.btn:not(.btn-rojo):not(.btn-sec)")
    exigir(p.evaluate("loteActual().pesos[0].kg") == 41.0, "con motivo, la N° 1 queda en 41.0")
    exigir(any(a["accion"] == "CORRECCIÓN DE PESO" and a["antes"] == "42.50 kg"
               for a in p.evaluate("jornada.bitacora")),
           "el valor anterior queda en la bitácora")

    p.click(".pesada >> nth=2")
    p.fill("#motivoCambio", "carcasa decomisada por el inspector")
    p.click(".gaveta button.btn-rojo")
    exigir(p.evaluate("loteActual().pesos.length") == 3, "la pesada anulada sigue en el registro")
    exigir(p.evaluate("vigentes(loteActual()).length") == 2, "pero no cuenta como vigente")
    exigir(p.evaluate("pesoLote(loteActual())") == 80.8, "ni suma: quedan 80.80 kg")
    exigir(p.inner_text(".correlativo .num") == "4",
           "el correlativo sigue en 4: un número anulado no se reutiliza")

    print("\n[7] Se termina el lote y recién ahí se pone el precio")
    p.click("text=Terminar lote")
    exigir("Lote culminado" in p.inner_text("h1"), "el lote queda culminado")
    exigir("80.80 kg" in p.inner_text(".paso-cierre"), "el paso 1 da el peso total: 80.80 kg")
    exigir(p.input_value("#precioLote") == "14.00", "propone el precio de referencia")
    p.fill("#precioLote", "15,50")
    p.dispatch_event("#precioLote", "change")
    exigir("S/ 1252.40" in p.inner_text("body"), "80.80 kg × S/ 15.50 = S/ 1252.40")
    p.screenshot(path=SALIDA / "4-culminado.png", full_page=True)

    print("\n[8] Lote de menudencias: por unidad, no por peso")
    p.click("text=Volver a la jornada")
    p.click("text=Nuevo lote de menudencias")
    p.select_option("#categoriaLote", "Menudencia de cordero")
    p.fill("#cantidadLote", "12")
    p.click("text=Abrir el lote")
    p.fill("#precioUnidad", "18.50")
    p.dispatch_event("#precioUnidad", "change")
    exigir("S/ 222.00" in p.inner_text("body"), "12 und × S/ 18.50 = S/ 222.00")
    p.click("button[aria-label='Sumar una']")
    exigir("S/ 240.50" in p.inner_text("body"), "sube a 13 und y recalcula")
    p.screenshot(path=SALIDA / "5-menudencias.png", full_page=True)
    p.click("text=Cerrar el lote")
    p.click("text=Volver a la jornada")

    print("\n[9] La jornada sobrevive a recargar el teléfono")
    antes = p.inner_text(".tiras")
    p.reload()
    exigir(p.inner_text(".tiras") == antes, "los totales siguen ahí tras recargar")
    exigir(p.locator(".lote").count() == 5, "los cinco lotes siguen ahí")
    exigir(p.evaluate("jornada.operador") == "Nilthon Chit", "y el responsable también")
    p.screenshot(path=SALIDA / "6-jornada-final.png", full_page=True)

    print("\n[10] La bitácora cuenta lo que pasó")
    p.click("text=Ver la bitácora")
    exigir(p.locator(".apunte.correccion").count() == 1, "la corrección aparece marcada")
    exigir(p.locator(".apunte.anulacion").count() == 1, "la anulación también")
    exigir("carcasa decomisada" in p.inner_text(".bitacora"), "con su motivo escrito")
    p.screenshot(path=SALIDA / "7-bitacora.png", full_page=True)
    p.click("button[aria-label='Volver']")

    print("\n[11] Los archivos de la jornada")
    p.click("text=Exportar: Excel, PDF y bitácora")
    with p.expect_download() as bajada:
        p.click("text=Descargar Excel (.xlsx)")
    libro = SALIDA / bajada.value.suggested_filename
    bajada.value.save_as(libro)
    with p.expect_download() as bajada:
        p.click("text=Descargar el acta en PDF")
    acta = SALIDA / bajada.value.suggested_filename
    bajada.value.save_as(acta)
    exigir(libro.suffix == ".xlsx" and libro.stat().st_size > 3000, f"baja {libro.name}")
    exigir(acta.suffix == ".pdf" and acta.stat().st_size > 2000, f"baja {acta.name}")
    p.screenshot(path=SALIDA / "8-exportar.png", full_page=True)

    ctx.close()
    nav.close()

# ---------------------------------------------------------------------------
# Auditoría de los archivos, ya fuera del navegador.
# ---------------------------------------------------------------------------
print("\n[12] El Excel, abierto con openpyxl")
from openpyxl import load_workbook

wb = load_workbook(libro)                      # con las fórmulas tal cual
exigir(wb.sheetnames == ["CONTROL", "PESOS POR LOTE", "DETALLE", "RESUMEN", "BITACORA"],
       f"cinco hojas en orden: {wb.sheetnames}")

lot = wb["PESOS POR LOTE"]
lotv = load_workbook(libro, data_only=True)["PESOS POR LOTE"]
etiquetas = [lot.cell(row=r, column=2).value for r in range(1, lot.max_row + 1)]
exigir(etiquetas.count("PESO TOTAL DEL LOTE") == 3,
       "un «peso total del lote» por cada lote de pesaje")
exigir(etiquetas.count("PRECIO TOTAL DEL LOTE") == 5,
       "y un «precio total del lote» por cada lote, menudencias incluidas")
exigir(any(str(lot.cell(row=r, column=3).value).startswith('=SUMIF(')
           for r in range(1, lot.max_row + 1)),
       "el peso total del bloque suma solo lo vigente, con fórmula")
exigir(any("ROUND(" in str(lot.cell(row=r, column=5).value)
           for r in range(1, lot.max_row + 1)),
       "y el precio total sale de peso por precio")
# El lote en curso no tiene precio puesto: sus celdas de dinero han de quedar
# vacías. Un cero ahí se leería como «cobrado a cero», que es otra cosa.
inicio_sin_precio = next(r for r in range(1, lot.max_row + 1)
                         if str(lot.cell(row=r, column=1).value or "").endswith("-L003"))
fin_sin_precio = next(r for r in range(inicio_sin_precio, lot.max_row + 1)
                      if lot.cell(row=r, column=2).value == "PRECIO TOTAL DEL LOTE")
dinero = [lotv.cell(row=r, column=5).value
          for r in range(inicio_sin_precio, fin_sin_precio + 1)
          if lotv.cell(row=r, column=5).value != "SUBTOTAL"]   # sin la cabecera
exigir(all(v is None for v in dinero),
       f"un lote sin precio deja las celdas de dinero vacías, no en cero: {dinero}")
exigir(lot.cell(row=lot.max_row, column=1).value == "TOTAL DE LA JORNADA",
       "el bloque final cierra con el total de la jornada")

exigir(abs(lotv.cell(row=lotv.max_row, column=3).value - 506.80) < 0.01,
       "el peso de la jornada cuadra: 506.80 kg")
exigir(abs(lotv.cell(row=lotv.max_row, column=5).value - 6385.70) < 0.01,
       "y el monto: S/ 6385.70")

det = wb["DETALLE"]
cab = [c.value for c in det[1]]
exigir(cab[0] == "ID_REGISTRO" and "ESTADO" in cab, "el detalle empieza por el identificador")
ids = [det.cell(row=r, column=1).value for r in range(2, det.max_row + 1)]
exigir(len(ids) == len(set(ids)), "no hay dos registros con el mismo identificador")
estados = [det.cell(row=r, column=12).value for r in range(2, det.max_row + 1)]
exigir(estados.count("ANULADO") == 1, "la anulada viaja al Excel marcada como tal")
exigir(det.cell(row=2, column=10).value.startswith("=IF("),
       "el subtotal es fórmula, no un número pegado")

res = wb["RESUMEN"]
formulas = [res.cell(row=r, column=7).value for r in range(2, res.max_row)]
exigir(any(str(f).startswith("=SUMIFS(DETALLE!") for f in formulas),
       "el peso del resumen se calcula contra el detalle")
exigir(any(str(res.cell(row=r, column=5).value).startswith("=COUNTIFS(DETALLE!")
           for r in range(2, res.max_row)),
       "y la cantidad también")

con = wb["CONTROL"]
textos = [str(c.value) for fila in con.iter_rows() for c in fila if c.value is not None]
exigir("REG-CAM-001" in textos and "Nilthon Chit" in textos,
       "el control documental lleva formato y responsable")
exigir(any(str(t).startswith('=COUNTIF(DETALLE!L:L,"ANULADO")') for t in textos),
       "y cuenta los anulados con fórmula")

# Y con los valores en frío, que es como lo lee un visor de teléfono.
wbv = load_workbook(libro, data_only=True)
resv = wbv["RESUMEN"]
total_frio = resv.cell(row=resv.max_row, column=9).value
exigir(abs(total_frio - 6385.70) < 0.01,
       f"el total en frío cuadra con la pantalla: S/ {total_frio}")

print("\n[13] El PDF, abierto con pypdf")
from pypdf import PdfReader

lector = PdfReader(str(acta))
texto_pdf = "\n".join(pagina.extract_text() or "" for pagina in lector.pages)
exigir(len(lector.pages) >= 1, f"{len(lector.pages)} página(s)")
exigir("REG-CAM-001" in texto_pdf, "trae el código del formato")
exigir("Nilthon Chit" in texto_pdf, "trae al responsable del registro")
exigir("ANULADO" in texto_pdf, "declara la pesada anulada")
exigir("carcasa decomisada por el inspector" in texto_pdf, "con su motivo")
exigir("Conformidad del cliente" in texto_pdf, "y deja las dos firmas")

print("\n--- errores de página:", errores or "ninguno")
print("--- fallos:", fallos or "ninguno")
sys.exit(1 if (errores or fallos) else 0)
