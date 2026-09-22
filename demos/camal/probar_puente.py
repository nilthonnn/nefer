"""El puente con Android, probado desde este lado.

Dentro del APK la app no descarga: entrega los bytes a Java en base64 y Android
los escribe en Descargas. Esa frontera es donde un archivo se corrompe sin que
nadie lo note —un byte mal codificado y el Excel abre roto en la oficina, tres
días después—, así que aquí se finge el puente, se recogen los bytes tal como
los recibiría Java, se decodifican y se abren con openpyxl y pypdf.

También se comprueba el botón de atrás del teléfono, que es JavaScript de este
lado aunque lo dispare Android.

    python3 demos/camal/probar_puente.py salidas/
"""

import base64
import pathlib
import sys

from playwright.sync_api import sync_playwright

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from _navegador import ruta_del_navegador

APP = pathlib.Path(__file__).resolve().parent / "app-camal.html"
SALIDA = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "salidas")
SALIDA.mkdir(parents=True, exist_ok=True)

# Lo que hace PuenteArchivos.java, escrito en JavaScript: recoge lo que llega
# y devuelve cadena vacía, que es su forma de decir «guardado».
PUENTE_FINGIDO = """
window.PuenteCamal = {
  recibido: [],
  guardar: function (nombre, base64, mime) {
    this.recibido.push({ nombre: nombre, base64: base64, mime: mime });
    return "";
  },
  version: function () { return "3.0"; }
};
"""

errores, fallos = [], []


def exigir(condicion, dicho):
    print(("  ok  " if condicion else "FALLA ") + dicho)
    if not condicion:
        fallos.append(dicho)


with sync_playwright() as pw:
    nav = pw.chromium.launch(executable_path=ruta_del_navegador())
    ctx = nav.new_context(viewport={"width": 393, "height": 851},
                          locale="es-PE", timezone_id="America/Lima")
    # El puente tiene que existir antes de que la app arranque, igual que
    # `addJavascriptInterface` antes de `loadUrl`.
    ctx.add_init_script(PUENTE_FINGIDO)
    p = ctx.new_page()
    p.on("pageerror", lambda e: errores.append(str(e)))
    p.goto(APP.as_uri())

    print("\n[1] La app ve el puente y deja de intentar descargar")
    exigir(p.evaluate("puente !== null"), "la app reconoce que corre dentro del APK")
    exigir(not p.evaluate("enMarco"), "y sabe que no está en un marco")

    print("\n[2] El Excel y el acta cruzan el puente")
    p.fill("#operador", "Nilthon Chit")
    p.dispatch_event("#operador", "change")
    p.click("text=Exportar: Excel, PDF y bitácora")
    p.click("text=Descargar Excel (.xlsx)")
    p.click("text=Descargar el acta en PDF")

    recibido = p.evaluate("window.PuenteCamal.recibido")
    exigir(len(recibido) == 2, f"llegaron {len(recibido)} archivos a Android")

    libro = next((r for r in recibido if r["nombre"].endswith(".xlsx")), None)
    acta = next((r for r in recibido if r["nombre"].endswith(".pdf")), None)
    exigir(libro is not None and acta is not None, "uno es el libro y el otro el acta")
    exigir(libro["mime"].endswith("spreadsheetml.sheet"),
           f"el libro va con su tipo: {libro['mime']}")
    exigir(acta["mime"] == "application/pdf", "y el acta con el suyo")
    exigir("Guardado en Descargas" in p.inner_text("#notaCopia")
           or "guardado en Descargas" in p.inner_text("#notaCopia"),
           "la app avisa de que Android lo guardó")

    print("\n[3] Los bytes llegan intactos, no parecidos")
    ruta_libro = SALIDA / libro["nombre"]
    ruta_acta = SALIDA / acta["nombre"]
    ruta_libro.write_bytes(base64.b64decode(libro["base64"]))
    ruta_acta.write_bytes(base64.b64decode(acta["base64"]))

    exigir(ruta_libro.read_bytes()[:2] == b"PK", "el .xlsx decodificado es un zip")
    exigir(ruta_acta.read_bytes()[:5] == b"%PDF-", "y el .pdf empieza por %PDF-")

    print("\n[4] Y los abre quien tiene que abrirlos")
    from openpyxl import load_workbook
    wb = load_workbook(ruta_libro)
    exigir(wb.sheetnames == ["CONTROL", "PESOS POR LOTE", "DETALLE", "RESUMEN", "BITACORA"],
           "el Excel que cruzó el puente tiene sus cinco hojas")

    from pypdf import PdfReader
    texto = "\n".join(pg.extract_text() or "" for pg in PdfReader(str(ruta_acta)).pages)
    exigir("REG-CAM-001" in texto and "Nilthon Chit" in texto,
           "y el acta trae su control documental")

    print("\n[5] El botón de atrás del teléfono")
    exigir(p.evaluate("atras()") is True, "desde exportar, vuelve a la jornada")
    exigir(p.evaluate("vista.pantalla") == "jornada", "y ahí está")
    exigir(p.evaluate("atras()") is False,
           "en la jornada contesta que no: ahí Android decide si cierra")

    p.click(".lote >> nth=0")
    p.click(".pesada >> nth=0")
    exigir(p.evaluate("atras()") is True, "con la gaveta abierta, la cierra")
    exigir(p.evaluate("document.getElementById('gaveta').innerHTML") == "",
           "y la gaveta queda cerrada, sin salir de la pantalla")

    ctx.close()
    nav.close()

print("\n--- errores de página:", errores or "ninguno")
print("--- fallos:", fallos or "ninguno")
sys.exit(1 if (errores or fallos) else 0)
