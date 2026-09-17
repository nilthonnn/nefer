"""De lo que hay en la oficina a fragmentos indexables.

Tres insumos, que son los tres que existen de verdad en un taller:

- el **historial de fallas**, una orden de trabajo por entrada: que fallo,
  cual resulto ser la causa y que se hizo. En JSON si alguien lo exporto asi,
  y en Excel —que es donde esta de verdad— con una fila por orden;
- los **manuales** del fabricante, en PDF, Word, Excel, Markdown o texto
  plano, troceados por seccion para que un procedimiento no se recupere
  partido a la mitad;
- las **actas** de despacho y recepcion que ya genera nefer —el manifiesto
  JSON o el propio Excel llenado—, de donde salen los componentes observados
  o dañados con su observacion escrita.

Lo ultimo es lo que hace que esto no sea un buscador de manuales mas: el
historial de la flota propia, con la observacion que el tecnico escribio a
mano, es la fuente que ningun manual OEM trae.

Que un Excel sea un historial, un acta o un manual no se pregunta: se mira.
Un acta de nefer trae su hoja REPORTE con el titulo del formato; un historial
trae columnas que se llaman falla, causa y solucion; lo demas es un manual.
"""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from . import documentos, texto as _texto
from .indice import Fragmento

EXTENSIONES_MANUAL = documentos.EXTENSIONES
EXTENSIONES = EXTENSIONES_MANUAL | {".json"}

CAMPOS_INFORME = ("codigo_ot", "fecha", "codigo_equipo", "modelo_equipo",
                  "categoria", "resumen_falla", "causa_raiz", "solucion_aplicada",
                  "horas_hombre", "horometro")


class ErrorIngesta(ValueError):
    """El documento no se puede convertir en fragmentos."""


def marca(clave: str) -> str:
    """Cuatro letras que identifican al archivo del que sale un fragmento.

    Los identificadores que se inventan —el de una seccion de manual, el de
    un informe sin numero de orden— salen del nombre del archivo, y dos
    carpetas distintas tienen manuales que se llaman igual: `2024/motor.md` y
    `2025/motor.md` daban el mismo identificador y uno de los dos desaparecia
    del indice en silencio, contado como indexado. Con esto no chocan.

    No se le pone a un codigo de orden de verdad: ese ya es unico, y que el
    mismo informe exportado dos veces caiga en el mismo sitio es lo correcto.
    """
    return hashlib.blake2b(str(clave).encode("utf-8"), digest_size=2).hexdigest()


def clave_de(ruta, raiz=None) -> str:
    """Con que nombre se identifica un documento, venga de la maquina que venga.

    La ruta completa no sirve: el mismo manual esta en `/home/ana/manuales`
    en la laptop de la oficina y en `/srv/nefer/manuales` en el servidor, y
    con la ruta dentro del identificador los dos son documentos distintos.
    Subir los dos a PostgreSQL duplicaba el historial entero en vez de
    actualizarlo, y reindexar despues de mover la carpeta rehacia todo.

    Lo que identifica al documento es su lugar **dentro del corpus**: la ruta
    relativa a la carpeta que se indexo. Eso sigue distinguiendo `2024/motor.md`
    de `2025/motor.md` —que es para lo que se invento la marca— y no cambia
    porque el corpus este colgado en otro sitio.

    Sin carpeta de referencia queda la carpeta que lo contiene mas el nombre
    —`2024/motor.md`—, que es lo minimo que sigue distinguiendo a los dos
    manuales del ejemplo sin volver a depender de donde este montado todo.
    """
    ruta = Path(ruta)
    if raiz is not None:
        try:
            return ruta.resolve().relative_to(Path(raiz).resolve()).as_posix()
        except ValueError:
            pass                      # no cuelga de ahi; queda el nombre corto
    padre = ruta.parent.name
    return f"{padre}/{ruta.name}" if padre else ruta.name


# ------------------------------------------------------------- historial

def de_informe(informe: dict, fuente: str = "", n: int = 0,
               clave: str = "") -> Fragmento:
    """Una orden de trabajo cerrada -> un fragmento."""
    if not isinstance(informe, dict):
        raise ErrorIngesta(f"{fuente}: cada informe debe ser un objeto JSON.")
    declarado = str(informe.get("codigo_ot") or "").strip()
    # Sin numero de orden hay que inventarle uno, y el inventado lleva la
    # marca del archivo: si no, el tercer informe sin OT de dos historiales
    # distintos seria el mismo fragmento.
    ot = declarado or f"SIN-OT-{marca(clave or fuente)}-{n + 1}"

    pasos = [str(p).strip() for p in informe.get("pasos") or [] if str(p).strip()]
    herramientas = [str(h).strip() for h in informe.get("herramientas") or [] if str(h).strip()]
    repuestos = [str(r).strip() for r in informe.get("repuestos") or [] if str(r).strip()]

    lineas = [f"ORDEN DE TRABAJO {ot}"]
    for etiqueta, clave in (("Equipo", "codigo_equipo"), ("Modelo", "modelo_equipo"),
                            ("Fecha", "fecha"), ("Falla reportada", "resumen_falla"),
                            ("Causa raiz confirmada", "causa_raiz"),
                            ("Solucion aplicada", "solucion_aplicada")):
        valor = str(informe.get(clave) or "").strip()
        if valor:
            lineas.append(f"{etiqueta}: {valor}")
    if pasos:
        lineas.append("Procedimiento: " + " ".join(f"{i}. {p}" for i, p in enumerate(pasos, 1)))
    if herramientas:
        lineas.append("Herramientas: " + ", ".join(herramientas))
    if repuestos:
        lineas.append("Repuestos: " + ", ".join(repuestos))
    otros = informe.get("otros_campos") or {}
    for clave, valor in otros.items():
        if str(valor).strip():
            lineas.append(f"{clave}: {valor}")
    cuerpo = "\n".join(lineas)

    declarados = informe.get("codigos_dtc") or informe.get("codigo_dtc") or []
    if isinstance(declarados, str):
        declarados = [declarados]
    codigos = [_texto.normalizar_dtc(c) for c in declarados if str(c).strip()]
    for codigo in _texto.codigos_dtc(cuerpo):
        if codigo not in codigos:
            codigos.append(codigo)

    metadatos = {clave: informe[clave] for clave in CAMPOS_INFORME if informe.get(clave)}
    for clave, valor in otros.items():
        # Con prefijo, para no pisar un campo del esquema con una columna que
        # se llamaba igual en el Excel de la oficina.
        metadatos.setdefault("col_" + (clave_columna(clave).replace(" ", "_") or "x"),
                             str(valor))
    metadatos.update({
        "codigo_ot": ot,
        "codigos_dtc": codigos,
        "pasos": pasos,
        "herramientas": herramientas,
        "repuestos": repuestos,
        "torques": _texto.torques(cuerpo),
    })
    return Fragmento(id=f"ot:{ot}", texto=cuerpo, fuente=fuente or ot,
                     tipo="informe", metadatos=metadatos)


def de_historial(datos, fuente: str = "", clave: str = "") -> list[Fragmento]:
    """Lista de ordenes de trabajo -> fragmentos."""
    informes = datos.get("informes", []) if isinstance(datos, dict) else datos
    if not isinstance(informes, list):
        raise ErrorIngesta(f"{fuente}: se esperaba una lista de informes.")
    return [de_informe(informe, fuente, n, clave)
            for n, informe in enumerate(informes)]


# Como se llama cada campo en el Excel del taller. Primero lo que se busca,
# luego todas las formas en que esta escrito en los libros de verdad.
COLUMNAS_INFORME = {
    "codigo_ot": ("ot", "n ot", "nro ot", "num ot", "numero ot", "orden",
                  "orden de trabajo", "codigo ot", "ot n", "n orden", "os"),
    "fecha": ("fecha", "fecha de atencion", "fecha atencion", "fecha cierre",
              "fecha de cierre", "dia"),
    "codigo_equipo": ("equipo", "codigo equipo", "cod equipo", "codigo de equipo",
                      "unidad", "maquina", "flota", "placa", "interno"),
    "modelo_equipo": ("modelo", "descripcion del equipo", "marca modelo"),
    "categoria": ("categoria", "familia", "tipo de equipo", "linea"),
    "codigos_dtc": ("dtc", "codigo de falla", "codigo falla", "cod falla",
                    "codigo dtc", "spn", "fmi", "codigo de error"),
    "resumen_falla": ("falla", "sintoma", "sintomas", "descripcion de la falla",
                      "problema", "reporte", "falla reportada", "descripcion",
                      "detalle de la falla"),
    "causa_raiz": ("causa", "causa raiz", "causa basica", "diagnostico",
                   "causa de la falla", "origen"),
    "solucion_aplicada": ("solucion", "solucion aplicada", "trabajo realizado",
                          "accion", "accion correctiva", "reparacion",
                          "actividad realizada", "correctivo"),
    "pasos": ("pasos", "procedimiento", "secuencia"),
    "herramientas": ("herramientas", "herramienta", "herramental"),
    "repuestos": ("repuestos", "repuesto", "materiales", "partes", "insumos"),
    "horas_hombre": ("horas hombre", "hh", "horas", "tiempo"),
    "horometro": ("horometro", "horas de equipo", "km", "kilometraje"),
}

CAMPOS_LISTA = ("pasos", "herramientas", "repuestos", "codigos_dtc")


def clave_columna(nombre: str) -> str:
    """El nombre de una columna, comparable: sin tildes, sin simbolos ni N°."""
    limpio = re.sub(r"[^a-z0-9]+", " ", _texto.normalizar(nombre))
    return re.sub(r"\s+", " ", limpio).strip()


def _mapa_de_columnas(encabezados) -> dict[str, str]:
    """Que columna del Excel alimenta que campo del informe.

    Primero se reparten las columnas que se llaman exactamente como uno de
    los sinonimos, y solo despues las que se parecen. Al reves, un campo que
    se prueba antes se lleva por delante una columna que era de otro: con los
    sinonimos de `codigos_dtc` mirando primero, la columna «Falla» casaba con
    «codigo de falla» y el sintoma desaparecia del informe convertido en un
    codigo de averia que nadie escribio.
    """
    mapa: dict[str, str] = {}
    usados: set[str] = set()
    normales = {col: clave_columna(col) for col in encabezados if col}

    for exacto in (True, False):
        for campo, sinonimos in COLUMNAS_INFORME.items():
            if campo in mapa:
                continue
            for columna, normal in normales.items():
                if columna in usados:
                    continue
                # En la vuelta de parecidos solo cuenta que la columna
                # CONTENGA el sinonimo: «falla reportada» es de la falla, y
                # «falla» a secas no es «codigo de falla».
                casa = (normal in sinonimos if exacto else
                        any(s in normal for s in sinonimos if len(s) > 2))
                if casa:
                    mapa[campo] = columna
                    usados.add(columna)
                    break
    return mapa


def _lista_de_celda(valor) -> list[str]:
    """Una celda con varias cosas dentro: separadas por ; o por salto de linea."""
    partes = re.split(r"[;\n·•]+|(?<=[a-z0-9])\s*,\s*(?=[A-Za-zÁÉÍÓÚÑ])", str(valor))
    return [p.strip(" .-") for p in partes if p.strip(" .-")]


def de_historial_xlsx(ruta: str | Path, hoja: str | None = None,
                      clave: str | None = None) -> list[Fragmento]:
    """El historial de fallas tal como esta en el taller: una fila por orden."""
    ruta = Path(ruta)
    try:
        registros = documentos.tabla(ruta, hoja=hoja)
    except documentos.ErrorDocumento as exc:
        raise ErrorIngesta(str(exc)) from exc
    if not registros:
        raise ErrorIngesta(f"{ruta.name}: la hoja esta vacia.")

    mapa = _mapa_de_columnas(registros[0].keys())
    if "resumen_falla" not in mapa and "causa_raiz" not in mapa:
        raise ErrorIngesta(
            f"{ruta.name}: no se reconoce ninguna columna de falla ni de causa "
            f"(columnas: {', '.join(list(registros[0])[:8])}). Renombre la "
            "columna de la falla a 'Falla' o la de la causa a 'Causa raiz'.")

    informes = []
    for registro in registros:
        informe = {}
        for campo, columna in mapa.items():
            valor = registro.get(columna, "")
            if not str(valor).strip():
                continue
            informe[campo] = _lista_de_celda(valor) if campo in CAMPOS_LISTA else valor
        # Lo que no se supo mapear no se tira: el tecnico que la escribio, el
        # sistema afectado o el costo son columnas por las que se busca.
        otros = {c: v for c, v in registro.items()
                 if c not in mapa.values() and str(v).strip()}
        if otros:
            informe["otros_campos"] = otros
        if any(informe.get(c) for c in ("resumen_falla", "causa_raiz",
                                        "solucion_aplicada")):
            informes.append(informe)
    if not informes:
        raise ErrorIngesta(f"{ruta.name}: ninguna fila trae falla, causa ni solucion.")
    clave = clave or clave_de(ruta)
    return [de_informe(informe, ruta.name, n, clave)
            for n, informe in enumerate(informes)]


# ----------------------------------------------------------------- actas

def de_manifiesto(manifiesto: dict, fuente: str = "",
                  clave: str = "") -> list[Fragmento]:
    """Un acta de nefer -> un fragmento por acta y uno por hallazgo."""
    enc = manifiesto.get("encabezado") or {}
    equipo = str(enc.get("codigo_equipo") or "").strip()
    acta = str(enc.get("n_acta") or "").strip()
    # Dos actas pueden llevar el mismo numero en dos archivos distintos: la
    # marca del archivo evita que la segunda pise a la primera.
    base = (acta or equipo or Path(fuente).stem or "acta") + "@" + marca(clave or fuente)
    tipo_doc = str(enc.get("tipo_documento") or "").strip()

    comun = {
        "codigo_equipo": equipo,
        "modelo_equipo": str(enc.get("modelo_equipo") or ""),
        "categoria": str(enc.get("categoria") or ""),
        "fecha": str(enc.get("fecha") or ""),
        "cliente": str(enc.get("cliente") or ""),
        "obra": str(enc.get("obra") or ""),
        "n_acta": acta,
        "tipo_documento": tipo_doc,
    }
    comun = {k: v for k, v in comun.items() if v}
    # El horometro es el unico dato de uso que existe en toda la flota: de el
    # sale el ritmo de horas por dia y el proximo servicio.
    if isinstance(enc.get("horometro"), (int, float)) and not isinstance(
            enc.get("horometro"), bool):
        comun["horometro"] = enc["horometro"]

    cabecera = (f"ACTA DE {tipo_doc or 'MOVIMIENTO'} {acta}".strip()
                + f"\nEquipo: {equipo} {comun.get('modelo_equipo', '')}".rstrip())
    if enc.get("horometro") not in (None, ""):
        cabecera += f"\nHorometro: {enc['horometro']}"
    if manifiesto.get("resumen_ejecutivo"):
        cabecera += f"\nResumen: {manifiesto['resumen_ejecutivo']}"

    hallazgos = [c for c in manifiesto.get("inspeccion_componentes") or []
                 if isinstance(c, dict) and c.get("estado") in ("OBS", "D")]
    if hallazgos:
        cabecera += "\nHallazgos: " + "; ".join(
            f"{c.get('item', '')} ({c.get('estado')}): {c.get('observacion', '')}".strip()
            for c in hallazgos)

    fragmentos = [Fragmento(
        id=f"acta:{base}", texto=cabecera, fuente=fuente or base, tipo="acta",
        metadatos={**comun, "codigos_dtc": _texto.codigos_dtc(cabecera),
                   "hallazgos": len(hallazgos)})]

    # Un consumible que no retorno, o que retorno mal, es un hecho registrado
    # del equipo igual que un componente dañado. Ademas es lo unico que
    # sobrevive al pasar por el Excel: la hoja de inspeccion no se relee.
    faltantes = [c for c in manifiesto.get("consumibles") or []
                 if isinstance(c, dict)
                 and str(c.get("estado_recepcion") or "OK").upper() not in ("OK", "")]
    for n, consumible in enumerate(faltantes, 1):
        descripcion = str(consumible.get("descripcion") or "").strip()
        estado = str(consumible.get("estado_recepcion") or "").strip()
        detalle = " ".join(str(consumible.get(clave) or "").strip()
                           for clave in ("texto_recepcion", "recuperacion")).strip()
        cuerpo = (f"CONSUMIBLE {descripcion}\nEstado en recepcion: {estado}\n"
                  f"{detalle}\nEquipo: {equipo} {comun.get('modelo_equipo', '')}").rstrip()
        fragmentos.append(Fragmento(
            id=f"acta:{base}:c{n}", texto=cuerpo, fuente=fuente or base, tipo="acta",
            metadatos={**comun, "item": descripcion, "estado": estado,
                       # Una barra de tierra que no volvio es una recuperacion
                       # que se factura, no una averia del equipo. Se indexa
                       # —se busca por ella— pero no cuenta como falla.
                       "clase": "consumible",
                       "codigos_dtc": _texto.codigos_dtc(cuerpo)}))

    # Cada componente observado o dañado es un antecedente por si mismo: la
    # observacion que el tecnico escribio en campo es el dato que se busca.
    for n, componente in enumerate(hallazgos, 1):
        item = str(componente.get("item") or "").strip()
        observacion = str(componente.get("observacion") or "").strip()
        cuerpo = (f"COMPONENTE {item}\nEstado: {componente.get('estado')}\n"
                  f"Observacion: {observacion}\n"
                  f"Equipo: {equipo} {comun.get('modelo_equipo', '')}".rstrip())
        fragmentos.append(Fragmento(
            id=f"acta:{base}:{n}", texto=cuerpo, fuente=fuente or base, tipo="acta",
            metadatos={**comun, "item": item, "estado": componente.get("estado"),
                       "clase": "componente",
                       "codigos_dtc": _texto.codigos_dtc(cuerpo)}))
    return fragmentos


def de_acta_xlsx(ruta: str | Path, clave: str | None = None) -> list[Fragmento]:
    """Un Excel de acta ya llenado -> fragmentos, pasando por el extractor."""
    from .. import extract

    ruta = Path(ruta)
    try:
        manifiesto = extract.extraer(ruta)
    except Exception as exc:
        raise ErrorIngesta(f"{ruta}: no se pudo leer como acta ({exc}).") from exc
    return de_manifiesto(manifiesto, fuente=ruta.name,
                         clave=clave or clave_de(ruta))


# -------------------------------------------------------------- manuales

_RE_TITULO = re.compile(r"^(#{1,6})\s*(.+?)\s*#*$")
_RE_HERRAMIENTAS = re.compile(r"^\s*(?:herramientas?|herramental)\s*:\s*(.+)$", re.I)


def de_manual(contenido: str, fuente: str = "", equipo: str = "",
              categoria: str = "", clave: str = "") -> list[Fragmento]:
    """Un manual en Markdown o texto plano -> un fragmento por seccion.

    Se corta por titulo, no por numero de caracteres: un torque recuperado sin
    el paso que lo pide no le sirve a nadie en campo.
    """
    fragmentos: list[Fragmento] = []
    for seccion, cuerpo in _secciones(contenido):
        for n, trozo in enumerate(_texto.trocear(cuerpo), 1):
            encabezado = f"{seccion}\n" if seccion else ""
            cuerpo_trozo = f"{encabezado}{trozo}"
            herramientas = []
            for linea in trozo.splitlines():
                m = _RE_HERRAMIENTAS.match(linea)
                if m:
                    herramientas.extend(
                        h.strip(" .-") for h in re.split(r"[,;]", m.group(1)) if h.strip(" .-"))
            fragmentos.append(Fragmento(
                id=(f"man:{Path(fuente).stem or 'manual'}"
                    f"@{marca(clave or fuente)}:{len(fragmentos) + 1}"),
                texto=cuerpo_trozo,
                fuente=fuente,
                tipo="manual",
                metadatos={
                    "seccion": seccion,
                    "codigo_equipo": equipo,
                    "categoria": categoria,
                    "codigos_dtc": _texto.codigos_dtc(cuerpo_trozo),
                    "torques": _texto.torques(cuerpo_trozo),
                    "herramientas": herramientas,
                    "parte": n,
                },
            ))
    return fragmentos


def _secciones(contenido: str) -> list[tuple[str, str]]:
    """Parte el documento por titulos Markdown; sin titulos, es una sola seccion."""
    secciones: list[tuple[str, list[str]]] = []
    actual = ("", [])
    for linea in str(contenido).splitlines():
        m = _RE_TITULO.match(linea)
        if m:
            if actual[1]:
                secciones.append(actual)
            actual = (m.group(2).strip(), [])
        else:
            actual[1].append(linea)
    if actual[1]:
        secciones.append(actual)
    return [(titulo, "\n".join(lineas).strip())
            for titulo, lineas in secciones if "\n".join(lineas).strip()]


# ------------------------------------------------------------- recorrido

def es_acta_de_nefer(ruta: Path) -> bool:
    """Un Excel generado por nefer trae su hoja REPORTE y el titulo del formato."""
    try:
        from openpyxl import load_workbook

        libro = load_workbook(ruta, read_only=True, data_only=True)
    except Exception:
        return False
    try:
        if not any(h.upper().startswith("REPORTE") for h in libro.sheetnames):
            return False
        hoja = libro.worksheets[0]
        for fila in hoja.iter_rows(min_row=1, max_row=3, values_only=True):
            for valor in fila:
                if isinstance(valor, str) and "REPORTE FOTOGR" in valor.upper():
                    return True
        return False
    finally:
        libro.close()


def de_excel(ruta: str | Path, clave: str | None = None) -> list[Fragmento]:
    """Un Excel puede ser tres cosas distintas; aqui se decide cual es.

    Se prueba de lo mas especifico a lo mas general: acta de nefer, historial
    de fallas con columnas reconocibles y, si no es ninguno, un manual con
    una seccion por hoja.
    """
    ruta = Path(ruta)
    clave = clave or clave_de(ruta)
    if es_acta_de_nefer(ruta):
        return de_acta_xlsx(ruta, clave)
    try:
        return de_historial_xlsx(ruta, clave=clave)
    except ErrorIngesta:
        return de_manual(documentos.leer(ruta), fuente=ruta.name, clave=clave)


def de_archivo(ruta: str | Path, raiz=None) -> list[Fragmento]:
    """Fragmentos de un archivo, deduciendo que es por su forma.

    `raiz` es la carpeta que se esta indexando: de ella sale el identificador
    de lo que no trae uno propio (ver `clave_de`).
    """
    ruta = Path(ruta)
    clave = clave_de(ruta, raiz)
    sufijo = ruta.suffix.lower()
    if sufijo == ".json":
        try:
            datos = json.loads(ruta.read_text(encoding="utf-8"))
        except json.JSONDecodeError as exc:
            raise ErrorIngesta(f"{ruta}: JSON ilegible ({exc}).") from exc
        if isinstance(datos, dict) and "encabezado" in datos:
            fragmentos = de_manifiesto(datos, fuente=ruta.name, clave=clave)
        elif isinstance(datos, list) or (isinstance(datos, dict)
                                         and "informes" in datos):
            fragmentos = de_historial(datos, fuente=ruta.name, clave=clave)
        else:
            raise ErrorIngesta(
                f"{ruta}: no parece ni un acta (falta 'encabezado') ni un "
                "historial (falta 'informes').")
    elif sufijo in (".xlsx", ".xlsm"):
        fragmentos = de_excel(ruta, clave)
    elif sufijo in EXTENSIONES_MANUAL:
        try:
            fragmentos = de_manual(documentos.leer(ruta), fuente=ruta.name,
                                   clave=clave)
        except documentos.ErrorDocumento as exc:
            raise ErrorIngesta(str(exc)) from exc
        if not fragmentos:
            raise ErrorIngesta(f"{ruta.name}: no se saco texto del documento.")
    else:
        raise ErrorIngesta(
            f"{ruta}: extension no soportada ({sufijo or 'sin extension'}).")

    # De donde salio cada fragmento, para poder reindexar solo lo que cambio.
    # Con guion bajo delante: los metadatos internos no entran en la busqueda.
    origen = str(ruta)
    for fragmento in fragmentos:
        fragmento.metadatos["_origen"] = origen
    return fragmentos


def es_indice(ruta: Path) -> bool:
    """Si el archivo es un indice de FixMate y no un documento.

    Pasa constantemente: el indice se guarda dentro de la misma carpeta que
    se indexa. Sin esto, la segunda pasada intenta leerse a si misma.
    """
    if ruta.suffix.lower() != ".json":
        return False
    try:
        with open(ruta, "rb") as fh:
            cabeza = fh.read(400)
    except OSError:
        return False
    # Las dos primeras llaves que escribe `Indice.guardar`. Los vectores
    # vienen mucho despues y un indice grande no cabe en una olfateada.
    return b'"version"' in cabeza and b'"embebedor"' in cabeza


def recorrer(rutas) -> list[Path]:
    """Archivos indexables de una lista de rutas, entrando en las carpetas."""
    return [archivo for _, archivo in recorrer_con_raiz(rutas)]


def recorrer_con_raiz(rutas) -> list[tuple[Path, Path]]:
    """Lo mismo, diciendo ademas de que carpeta salio cada archivo.

    De esa carpeta sale el identificador de los fragmentos (ver `clave_de`),
    asi que hay que llevarla hasta la ingesta y no perderla por el camino.
    Para un archivo suelto la referencia es la carpeta que lo contiene.
    """
    encontrados: list[tuple[Path, Path]] = []
    for ruta in rutas:
        ruta = Path(ruta)
        if ruta.is_dir():
            encontrados.extend(
                (ruta, a) for a in sorted(ruta.rglob("*"))
                if a.is_file() and a.suffix.lower() in EXTENSIONES
                and not a.name.startswith(".") and not es_indice(a))
        elif ruta.is_file():
            if es_indice(ruta):
                raise ErrorIngesta(
                    f"{ruta}: es un indice de FixMate, no un documento. Indexe "
                    "la carpeta con los manuales y el historial.")
            encontrados.append((ruta.parent, ruta))
        else:
            raise ErrorIngesta(f"no existe: {ruta}")
    return encontrados
