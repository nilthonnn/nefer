"""Leer lo que hay en la oficina: PDF, Word, Excel y texto plano.

El manual del fabricante no esta en Markdown. Esta en un PDF de 400 paginas,
en un Word que alguien escribio sobre la plantilla de la empresa y en un Excel
con una hoja por sistema. El historial de fallas tampoco esta en JSON: esta en
un Excel con una fila por orden de trabajo.

Aqui se convierte todo eso en texto, conservando lo unico que hace falta
conservar: donde empieza y termina cada seccion. Los titulos de Word y los
nombres de hoja de Excel salen como titulos Markdown, y asi el troceado por
secciones de `ingesta.py` funciona igual venga de donde venga el documento.
"""

from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from . import pdf_texto

EXTENSIONES_TEXTO = {".md", ".markdown", ".txt"}
EXTENSIONES_OFIMATICA = {".pdf", ".docx", ".xlsx", ".xlsm"}
# Formatos de hace veinte años que siguen en la carpeta compartida.
EXTENSIONES_HEREDADAS = {".doc": ".docx", ".xls": ".xlsx", ".odt": ".docx",
                         ".ods": ".xlsx", ".rtf": ".docx"}
EXTENSIONES = EXTENSIONES_TEXTO | EXTENSIONES_OFIMATICA | set(EXTENSIONES_HEREDADAS)

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class ErrorDocumento(ValueError):
    """El documento no se pudo leer."""


def leer(ruta: str | Path) -> str:
    """Texto de un documento, sea cual sea su formato."""
    ruta = Path(ruta)
    if not ruta.is_file():
        raise ErrorDocumento(f"no existe el documento {ruta}")
    sufijo = ruta.suffix.lower()

    if sufijo in EXTENSIONES_TEXTO:
        return ruta.read_text(encoding="utf-8", errors="replace")
    if sufijo == ".pdf":
        try:
            return pdf_texto.extraer(ruta)
        except pdf_texto.ErrorPDFTexto as exc:
            raise ErrorDocumento(str(exc)) from exc
    if sufijo == ".docx":
        return leer_docx(ruta)
    if sufijo in (".xlsx", ".xlsm"):
        return leer_xlsx(ruta)
    if sufijo in EXTENSIONES_HEREDADAS:
        with _convertido(ruta) as moderno:
            return leer(moderno)
    raise ErrorDocumento(f"{ruta.name}: formato no soportado ({sufijo or 'sin extension'}).")


# ------------------------------------------------------------------ Word

def leer_docx(ruta: str | Path) -> str:
    """Texto de un .docx. Los titulos de Word salen como titulos Markdown.

    Se abre con `zipfile` y se lee el XML: un .docx es un zip con un XML
    dentro, y no hace falta una dependencia para eso.
    """
    ruta = Path(ruta)
    try:
        with zipfile.ZipFile(ruta) as z:
            crudo = z.read("word/document.xml")
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ErrorDocumento(
            f"{ruta.name}: no es un .docx legible ({exc}). Si es un .doc de los "
            "antiguos, cambiele la extension o abralo y guardelo como .docx."
        ) from exc

    try:
        raiz = ET.fromstring(crudo)
    except ET.ParseError as exc:
        raise ErrorDocumento(f"{ruta.name}: XML ilegible ({exc}).") from exc

    cuerpo = raiz.find(f"{_W}body")
    lineas: list[str] = []
    for nodo in list(cuerpo if cuerpo is not None else raiz):
        etiqueta = nodo.tag.replace(_W, "")
        if etiqueta == "p":
            lineas.append(_parrafo_docx(nodo))
        elif etiqueta == "tbl":
            lineas.extend(_tabla_docx(nodo))
    return _juntar(lineas)


def _parrafo_docx(nodo) -> str:
    texto = _texto_docx(nodo)
    estilo = nodo.find(f"{_W}pPr/{_W}pStyle")
    nombre = (estilo.get(f"{_W}val") if estilo is not None else "") or ""
    m = re.fullmatch(r"(?:Heading|Ttulo|Titulo|Ttulo)\s*(\d)", nombre, re.I)
    if m and texto.strip():
        # Un titulo de Word es una seccion; asi lo trocea la ingesta.
        return "#" * min(int(m.group(1)), 6) + " " + texto.strip()
    if nombre.lower().startswith(("heading", "titulo", "título")) and texto.strip():
        return "## " + texto.strip()
    return texto


def _texto_docx(nodo) -> str:
    partes = []
    for hijo in nodo.iter():
        etiqueta = hijo.tag.replace(_W, "")
        if etiqueta == "t":
            partes.append(hijo.text or "")
        elif etiqueta == "tab":
            partes.append("\t")
        elif etiqueta in ("br", "cr"):
            partes.append("\n")
    return "".join(partes).strip()


def _tabla_docx(tabla) -> list[str]:
    filas = []
    for fila in tabla.findall(f"{_W}tr"):
        celdas = [re.sub(r"\s+", " ", _texto_docx(c)).strip()
                  for c in fila.findall(f"{_W}tc")]
        if any(celdas):
            filas.append(" | ".join(celdas))
    return filas


# ----------------------------------------------------------------- Excel

def leer_xlsx(ruta: str | Path, maximo_filas: int = 5000) -> str:
    """Texto de un libro de Excel: una seccion por hoja, una linea por fila."""
    hojas = filas_xlsx(ruta, maximo_filas=maximo_filas)
    lineas: list[str] = []
    for hoja, filas in hojas.items():
        lineas.append(f"## {hoja}")
        for fila in filas:
            texto = " | ".join(_celda(v) for v in fila).strip(" |")
            if texto.strip():
                lineas.append(texto)
        lineas.append("")
    return _juntar(lineas)


def filas_xlsx(ruta: str | Path, maximo_filas: int = 5000) -> dict[str, list[list]]:
    """Las celdas de cada hoja, sin formato. Las formulas, ya calculadas."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - openpyxl es dependencia
        raise ErrorDocumento("leer Excel necesita openpyxl.") from exc

    ruta = Path(ruta)
    try:
        # data_only: interesa el valor que se ve, no la formula que lo calcula.
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except Exception as exc:
        raise ErrorDocumento(f"{ruta.name}: no se pudo abrir como Excel ({exc}).") from exc

    try:
        hojas = {}
        for hoja in libro.worksheets:
            filas = []
            for n, fila in enumerate(hoja.iter_rows(values_only=True)):
                if n >= maximo_filas:
                    break
                if any(v is not None and str(v).strip() for v in fila):
                    filas.append(list(fila))
            if filas:
                hojas[hoja.title] = filas
        return hojas
    finally:
        libro.close()


def _celda(valor) -> str:
    if valor is None:
        return ""
    if hasattr(valor, "strftime"):
        return valor.strftime("%Y-%m-%d")
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return re.sub(r"\s+", " ", str(valor)).strip()


def tabla(ruta: str | Path, hoja: str | None = None) -> list[dict[str, str]]:
    """Una hoja de Excel como lista de diccionarios, con su fila de encabezado.

    El encabezado no siempre esta en la primera fila —arriba suele haber un
    logo, un titulo y dos filas en blanco—, asi que se busca: la fila con mas
    celdas de texto distintas, entre las diez primeras, es el encabezado.
    """
    hojas = filas_xlsx(ruta)
    if not hojas:
        return []
    nombre = hoja or next(iter(hojas))
    if nombre not in hojas:
        raise ErrorDocumento(
            f"{Path(ruta).name}: no tiene una hoja llamada {nombre!r} "
            f"(tiene {', '.join(hojas)}).")
    filas = hojas[nombre]

    mejor, puntaje = 0, 0
    for i, fila in enumerate(filas[:10]):
        textos = {_celda(v).lower() for v in fila if _celda(v)}
        # Una fila de datos repite valores y trae numeros; una de encabezado
        # es texto y no se repite.
        cuenta = len([v for v in fila if _celda(v) and not _es_numero(v)])
        if cuenta > puntaje and len(textos) == cuenta:
            mejor, puntaje = i, cuenta
    encabezados = [_celda(v) for v in filas[mejor]]

    registros = []
    for fila in filas[mejor + 1:]:
        registro = {}
        for clave, valor in zip(encabezados, fila):
            if clave:
                registro[clave] = _celda(valor)
        if any(registro.values()):
            registros.append(registro)
    return registros


def _es_numero(valor) -> bool:
    if isinstance(valor, bool):
        return False
    if isinstance(valor, (int, float)):
        return True
    try:
        float(str(valor).replace(",", "."))
        return True
    except (TypeError, ValueError):
        return False


# ------------------------------------------------------ formatos antiguos

class _convertido:
    """Convierte .doc/.xls/.odt a su formato moderno con LibreOffice."""

    def __init__(self, ruta: Path):
        self.ruta = Path(ruta)
        self.temporal = None

    def __enter__(self) -> Path:
        destino = EXTENSIONES_HEREDADAS[self.ruta.suffix.lower()]
        binario = shutil.which("soffice") or shutil.which("libreoffice")
        if not binario:
            raise ErrorDocumento(
                f"{self.ruta.name}: para leer un {self.ruta.suffix} hace falta "
                "LibreOffice (sudo apt-get install libreoffice), o guardelo "
                f"como {destino} desde Office.")
        self.temporal = tempfile.TemporaryDirectory(prefix="fixmate-doc-")
        tmp = self.temporal.name
        cmd = [binario, "--headless", "--norestore",
               f"-env:UserInstallation=file://{tmp}/perfil",
               "--convert-to", destino.lstrip("."), "--outdir", tmp, str(self.ruta)]
        try:
            proceso = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
        except subprocess.TimeoutExpired as exc:
            self.__exit__(None, None, None)
            raise ErrorDocumento(
                f"{self.ruta.name}: LibreOffice no respondio en 300 s.") from exc
        generado = Path(tmp) / (self.ruta.stem + destino)
        if not generado.exists():
            self.__exit__(None, None, None)
            raise ErrorDocumento(
                f"{self.ruta.name}: LibreOffice no pudo convertirlo "
                f"({(proceso.stderr or proceso.stdout).strip()[:160]}).")
        return generado

    def __exit__(self, *_):
        if self.temporal is not None:
            self.temporal.cleanup()
            self.temporal = None
        return False


def _juntar(lineas: list[str]) -> str:
    """Une las lineas dejando una sola linea en blanco entre parrafos."""
    salida: list[str] = []
    for linea in lineas:
        limpia = linea.rstrip()
        if not limpia.strip():
            if salida and salida[-1] != "":
                salida.append("")
        else:
            salida.append(limpia)
    return "\n".join(salida).strip() + "\n" if salida else ""
