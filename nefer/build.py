"""Genera el libro Excel del reporte a partir de un manifiesto JSON."""

from __future__ import annotations

import datetime as _dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.pagebreak import Break

from . import estilos as st
from . import layout, schema, textos

FORMATOS_IMAGEN = {".png", ".jpg", ".jpeg", ".gif", ".bmp"}
MARGEN_IMAGEN_PX = 4


# --------------------------------------------------------------------------- #
# Imagenes
# --------------------------------------------------------------------------- #
def _medir(ruta: Path) -> tuple[int, int]:
    from PIL import Image  # importado aqui para no exigir Pillow en `validar`

    with Image.open(ruta) as im:
        return im.size


def insertar_imagen(ws, ruta: Path, columna: str, fila: int,
                    ancho_px: int, alto_px: int) -> bool:
    """Coloca la imagen centrada dentro del bloque, respetando su relacion de aspecto.

    Devuelve False si el archivo no es un formato que Excel sepa incrustar.
    """
    if ruta.suffix.lower() not in FORMATOS_IMAGEN:
        return False
    try:
        origen_w, origen_h = _medir(ruta)
    except Exception:
        return False
    if origen_w <= 0 or origen_h <= 0:
        return False

    disponible_w = ancho_px - 2 * MARGEN_IMAGEN_PX
    disponible_h = alto_px - 2 * MARGEN_IMAGEN_PX
    escala = min(disponible_w / origen_w, disponible_h / origen_h)
    destino_w = max(1, int(origen_w * escala))
    destino_h = max(1, int(origen_h * escala))

    offset_x = (ancho_px - destino_w) // 2
    offset_y = (alto_px - destino_h) // 2

    img = XLImage(str(ruta))
    img.anchor = OneCellAnchor(
        _from=AnchorMarker(
            col=column_index_from_string(columna) - 1,
            colOff=pixels_to_EMU(offset_x),
            row=fila - 1,
            rowOff=pixels_to_EMU(offset_y),
        ),
        ext=XDRPositiveSize2D(pixels_to_EMU(destino_w), pixels_to_EMU(destino_h)),
    )
    ws.add_image(img)
    return True


# --------------------------------------------------------------------------- #
# Hoja principal: reporte fotografico
# --------------------------------------------------------------------------- #
def _dimensionar(ws, ultima_fila: int, n_bloques_foto: int, n_consumibles: int,
                 bandas=None, n_vistas: int = 0) -> None:
    for col, ancho in layout.ANCHOS_COLUMNA.items():
        ws.column_dimensions[col].width = ancho

    for fila in range(1, ultima_fila + 1):
        ws.row_dimensions[fila].height = layout.ALTO_FILA_ESTANDAR
    ws.row_dimensions[layout.FILA_SEPARADOR].height = layout.ALTO_FILA_SEPARADOR

    for i in range(n_bloques_foto):
        ws.row_dimensions[layout.bloque_foto(i)["fila_rotulo"]].height = layout.ALTO_FILA_ROTULO
    if n_consumibles:
        ws.row_dimensions[
            layout.fila_titulo_observaciones(n_bloques_foto, n_vistas)].height = \
            layout.ALTO_TITULO_SECCION_PT
    for j in range(n_consumibles):
        bloque = layout.bloque_consumible(j, n_bloques_foto, bandas, n_vistas)
        ws.row_dimensions[bloque["fila_encabezado"]].height = layout.ALTO_FILA_ROTULO
        ws.row_dimensions[bloque["fila_rotulo"]].height = layout.ALTO_FILA_TEXTO
        for fila in bloque["filas_recuperacion"]:
            ws.row_dimensions[fila].height = layout.ALTO_FILA_ROTULO


def _ruta_logo(enc: dict, raiz: Path) -> Path | None:
    """Logo de la organizacion, si el manifiesto declara uno."""
    return raiz / enc["logo"] if enc.get("logo") else None


def _cabecera(ws, enc: dict, raiz: Path, avisos: list[str]) -> None:
    logo = _ruta_logo(enc, raiz)
    st.escribir(ws, "A1:F3", None)
    if logo is not None:
        ancho_logo = sum(round(layout.ANCHOS_COLUMNA[c] * 7) + 5 for c in "ABCDEF")
        if not logo.exists():
            avisos.append(f"logo: no existe {logo}; el acta sale sin logo.")
        elif not insertar_imagen(ws, logo, "A", 1, ancho_logo, layout.alto_bloque_px(3)):
            avisos.append(f"logo: no se pudo incrustar {logo.name}; formato no soportado.")

    st.escribir(ws, "G1:R3", layout.TITULO_FORMATO,
                fuente=st.FUENTE_TITULO, alineacion=st.CENTRO_AJUSTADO,
                fill=st.FILL_CABECERA)
    control = (
        ("S1:Z1", "CÓDIGO", enc.get("codigo_formato", layout.CODIGO_FORMATO)),
        ("S2:Z2", "VERSIÓN", enc.get("version_formato", layout.VERSION_FORMATO)),
        ("S3:Z3", "FECHA", enc.get("fecha_formato", layout.FECHA_FORMATO)),
    )
    for rango, etiqueta, valor in control:
        texto = f"{etiqueta}: {valor}" if str(valor).strip() else f"{etiqueta}:"
        st.escribir(ws, rango, texto, fuente=st.FUENTE_META, alineacion=st.IZQUIERDA)

    etiquetas = [
        ("A4:F4", "N° ACTA:", "G4:Z4", enc.get("n_acta", "")),
        ("A5:F5", "N° GUÍA:", "G5:Z5", enc.get("n_guia", "")),
        ("A7:F7", "CLIENTE:", "G7:Z7", enc.get("cliente", "")),
        ("A8:F8", "EQUIPO:", "G8:Z8", enc.get("modelo_equipo", "")),
    ]
    for rango_et, texto, rango_val, valor in etiquetas:
        st.escribir(ws, rango_et, texto, fuente=st.FUENTE_ETIQUETA, fill=st.FILL_CABECERA)
        st.escribir(ws, rango_val, valor, alineacion=st.IZQUIERDA_AJUSTADA)

    # Fila 6: fecha + casillas de tipo de documento.
    st.escribir(ws, "A6:F6", "FECHA:", fuente=st.FUENTE_ETIQUETA, fill=st.FILL_CABECERA)
    fecha = _dt.date.fromisoformat(enc["fecha"])
    st.escribir(ws, "G6:M6", fecha, formato="DD/MM/YYYY")
    es_despacho = enc.get("tipo_documento") == "DESPACHO"
    st.escribir(ws, "N6:R6", "DESPACHO", fuente=st.FUENTE_META)
    st.escribir(ws, "S6:T6", "X" if es_despacho else None, fuente=st.FUENTE_META)
    st.escribir(ws, "U6:Y6", "RECEPCIÓN", fuente=st.FUENTE_META)
    st.escribir(ws, "Z6", "X" if not es_despacho else None, fuente=st.FUENTE_META)

    # Fila 9: codigo interno + horometro.
    st.escribir(ws, "A9:F9", "CÓDIGO:", fuente=st.FUENTE_ETIQUETA, fill=st.FILL_CABECERA)
    st.escribir(ws, "G9:L9", enc.get("codigo_equipo", ""))
    st.escribir(ws, "M9:R9", "HORÓMETRO:", fuente=st.FUENTE_ETIQUETA, fill=st.FILL_CABECERA)
    horometro = enc.get("horometro")
    if horometro == schema.REVISION_MANUAL:
        st.escribir(ws, "S9:Z9", schema.REVISION_MANUAL,
                    fuente=st.FUENTE_ROTULO, alineacion=st.CENTRO_AJUSTADO,
                    fill=st.FILL_RECUPERACION)
    else:
        st.escribir(ws, "S9:Z9", horometro, formato="0.0")


def _rejilla_fotografica(ws, fotos: list[dict], raiz: Path, avisos: list[str],
                         pintar: bool = True) -> int:
    if not pintar:
        return 0
    n_bloques = (len(fotos) + 1) // 2
    ancho = {layout.PANEL_IZQ[0]: layout.ancho_panel_px(layout.PANEL_IZQ),
             layout.PANEL_DER[0]: layout.ancho_panel_px(layout.PANEL_DER)}
    alto = layout.alto_bloque_px()

    for pos in range(n_bloques * 2):
        bloque = layout.bloque_foto(pos // 2)
        panel = layout.PANEL_IZQ if pos % 2 == 0 else layout.PANEL_DER
        col_ini, col_fin = panel
        r0, r1 = bloque["fila_imagen_inicio"], bloque["fila_imagen_fin"]
        rotulo = bloque["fila_rotulo"]

        st.escribir(ws, f"{col_ini}{r0}:{col_fin}{r1}", None)
        foto = fotos[pos] if pos < len(fotos) else None
        st.escribir(ws, f"{col_ini}{rotulo}:{col_fin}{rotulo}",
                    (foto or {}).get("descripcion", ""),
                    fuente=st.FUENTE_ROTULO, alineacion=st.CENTRO_AJUSTADO,
                    fill=st.FILL_CABECERA)

        if foto and foto.get("archivo"):
            ruta = raiz / foto["archivo"]
            if not insertar_imagen(ws, ruta, col_ini, r0, ancho[col_ini], alto):
                avisos.append(
                    f"foto_id {foto.get('foto_id')}: no se pudo incrustar {ruta.name} "
                    "(formato no soportado por Excel); el bloque queda en blanco."
                )
    return n_bloques


def bandas_de_cierre(consumibles: list[dict], tipo_documento: str) -> list[int]:
    """Cuantas franjas de cierre ocupa cada bloque de observacion."""
    return [max(1, len(textos.cierres_consumible(c, tipo_documento)))
            for c in consumibles]


def _bloques_consumibles(ws, consumibles: list[dict], n_bloques_foto: int,
                         raiz: Path, avisos: list[str],
                         tipo_documento: str = "RECEPCION", bandas=None,
                         n_vistas: int = 0) -> None:
    if not consumibles:
        return
    fila_titulo = layout.fila_titulo_observaciones(n_bloques_foto, n_vistas)
    st.escribir(ws, f"A{fila_titulo}:Z{fila_titulo}", "OBSERVACIONES",
                fuente=st.FUENTE_META)

    izq, der = layout.PANEL_IZQ, layout.PANEL_DER
    ancho = {izq[0]: layout.ancho_panel_px(izq), der[0]: layout.ancho_panel_px(der)}
    alto = layout.alto_bloque_px()

    # Dos contadores independientes: uno para lo que se recupera y otro para lo
    # que vuelve conforme. Cada franja se numera dentro de su propia serie.
    n_recuperacion = n_conforme = 0

    for j, cons in enumerate(consumibles):
        b = layout.bloque_consumible(j, n_bloques_foto, bandas, n_vistas)
        for panel, titulo, clave_foto, clave_texto in (
            (izq, "DESPACHO", "foto_despacho", "texto_despacho"),
            (der, "RECEPCIÓN", "foto_recepcion", "texto_recepcion"),
        ):
            c0, c1 = panel
            st.escribir(ws, f"{c0}{b['fila_encabezado']}:{c1}{b['fila_encabezado']}",
                        titulo, fuente=st.FUENTE_ROTULO, alineacion=st.CENTRO_AJUSTADO,
                        fill=st.FILL_CABECERA)
            st.escribir(ws, f"{c0}{b['fila_imagen_inicio']}:{c1}{b['fila_imagen_fin']}", None)
            st.escribir(ws, f"{c0}{b['fila_rotulo']}:{c1}{b['fila_rotulo']}",
                        cons.get(clave_texto)
                        or textos.texto_consumible(cons, titulo, tipo_documento),
                        fuente=st.FUENTE_ROTULO, alineacion=st.CENTRO_AJUSTADO,
                        fill=st.FILL_CABECERA)
            if cons.get(clave_foto):
                ruta = raiz / cons[clave_foto]
                if not insertar_imagen(ws, ruta, c0, b["fila_imagen_inicio"],
                                       ancho[c0], alto):
                    avisos.append(
                        f"consumible {j + 1} ({titulo}): no se pudo incrustar {ruta.name}."
                    )

        cierres = textos.cierres_consumible(cons, tipo_documento)
        for k, fila_rec in enumerate(b["filas_recuperacion"]):
            leyenda = ""
            if k < len(cierres):
                if cierres[k]["recupera"]:
                    n_recuperacion += 1
                    numero = n_recuperacion
                else:
                    n_conforme += 1
                    numero = n_conforme
                leyenda = textos.texto_cierre(cons, cierres[k], numero, k)
            # Sin nada que declarar la franja va vacia y en blanco: una banda
            # amarilla sin texto se lee como un dato que falta.
            st.escribir(ws, f"A{fila_rec}:Z{fila_rec}", leyenda,
                        fuente=st.FUENTE_ROTULO, alineacion=st.CENTRO_AJUSTADO,
                        fill=st.FILL_RECUPERACION if leyenda else None)


def piezas_del_acta(n_bloques_foto: int, bandas: list[int], plan: dict,
                    n_danos: int = 0, n_vistas: int = 0) -> list[dict]:
    """El acta como una pila de piezas que no se pueden partir.

    Cada pieza lleva su alto en puntos de la rejilla y la ultima fila que
    ocupa. Con eso `layout.reparto` dice en que hoja cae cada una, y el mismo
    reparto vale para el Excel y para el PDF: los dos entregables tienen que
    poder cotejarse hoja contra hoja.
    """
    piezas = [{"alto": layout.ALTO_CABECERA_PT,
               "fila_fin": layout.FILA_SEPARADOR}]
    # El informe fotografico: emparejado por vista, o la rejilla del formato.
    for i in range(n_vistas):
        piezas.append({"alto": layout.alto_bloque_pareado_pt(1),
                       "fila_fin": layout.bloque_pareado(i, layout.FILA_VISTAS)["fila_pie"]})
    for i in range(n_bloques_foto):
        piezas.append({"alto": layout.ALTO_BLOQUE_FOTO_PT,
                       "fila_fin": layout.bloque_foto(i)["fila_rotulo"]})

    if bandas:
        piezas.append({"alto": layout.ALTO_TITULO_SECCION_PT,
                       "fila_fin": layout.fila_titulo_observaciones(n_bloques_foto,
                                                                   n_vistas),
                       "abre_pagina": True, "arrastra": True})
        for j, n in enumerate(bandas):
            b = layout.bloque_consumible(j, n_bloques_foto, bandas, n_vistas)
            piezas.append({"alto": layout.alto_bloque_pareado_pt(n),
                           "fila_fin": b["filas_recuperacion"][-1]})

    titulo = plan.get("fila_danos")
    if titulo and n_danos:
        piezas.append({"alto": layout.ALTO_TITULO_SECCION_PT, "fila_fin": titulo,
                       "abre_pagina": True, "arrastra": True})
        for j in range(n_danos):
            piezas.append({"alto": layout.alto_bloque_pareado_pt(1),
                           "fila_fin": layout.bloque_pareado(j, titulo)["fila_pie"]})
    return piezas


def _saltos_de_pagina(ws, n_bloques_foto: int, bandas: list[int],
                      plan: dict, n_danos: int = 0, n_vistas: int = 0) -> None:
    """Fuerza los cortes de pagina para que ningun bloque quede partido."""
    piezas = piezas_del_acta(n_bloques_foto, bandas, plan, n_danos, n_vistas)
    hojas = layout.reparto(piezas)
    for i in range(len(piezas) - 1):
        if hojas[i + 1] != hojas[i]:
            ws.row_breaks.append(Break(id=piezas[i]["fila_fin"]))


def _seccion_pareada(ws, titulo: str | None, fila_titulo: int, entradas: list[dict],
                     raiz: Path, avisos: list[str], resaltar_pie: bool = False) -> None:
    """Imprime una seccion de bloques «lo que salio / lo que volvio».

    Cada entrada es {rotulo, foto_izq, foto_der, texto_izq, texto_der, pie}.
    Sin `titulo` no se escribe cabecera de seccion: es el caso del informe
    fotografico emparejado, que abre el acta y no lleva rotulo propio.
    """
    if not entradas:
        return
    if titulo:
        st.escribir(ws, f"A{fila_titulo}:Z{fila_titulo}", titulo, fuente=st.FUENTE_META)
        ws.row_dimensions[fila_titulo].height = layout.ALTO_TITULO_SECCION_PT

    izq, der = layout.PANEL_IZQ, layout.PANEL_DER
    ancho = {izq[0]: layout.ancho_panel_px(izq), der[0]: layout.ancho_panel_px(der)}
    alto = layout.alto_bloque_px()

    for i, entrada in enumerate(entradas):
        b = layout.bloque_pareado(i, fila_titulo)
        for clave in ("fila_encabezado", "fila_pie"):
            ws.row_dimensions[b[clave]].height = layout.ALTO_FILA_ROTULO
        ws.row_dimensions[b["fila_rotulo"]].height = layout.ALTO_FILA_TEXTO
        for panel, cabecera, clave_foto, clave_texto in (
            (izq, "DESPACHO", "foto_izq", "texto_izq"),
            (der, "RECEPCIÓN", "foto_der", "texto_der"),
        ):
            c0, c1 = panel
            st.escribir(ws, f"{c0}{b['fila_encabezado']}:{c1}{b['fila_encabezado']}",
                        cabecera, fuente=st.FUENTE_ROTULO,
                        alineacion=st.CENTRO_AJUSTADO, fill=st.FILL_CABECERA)
            st.escribir(ws, f"{c0}{b['fila_imagen_inicio']}:{c1}{b['fila_imagen_fin']}", None)
            st.escribir(ws, f"{c0}{b['fila_rotulo']}:{c1}{b['fila_rotulo']}",
                        entrada.get(clave_texto) or entrada.get("rotulo", ""),
                        fuente=st.FUENTE_ROTULO, alineacion=st.CENTRO_AJUSTADO,
                        fill=st.FILL_CABECERA)
            archivo = entrada.get(clave_foto)
            if archivo:
                ruta = raiz / archivo
                if not insertar_imagen(ws, ruta, c0, b["fila_imagen_inicio"],
                                       ancho[c0], alto):
                    avisos.append(
                        f"{titulo.lower()}: no se pudo incrustar {ruta.name}.")

        pie = entrada.get("pie") or ""
        st.escribir(ws, f"A{b['fila_pie']}:Z{b['fila_pie']}", pie,
                    fuente=st.FUENTE_ROTULO, alineacion=st.CENTRO_AJUSTADO,
                    fill=st.FILL_RECUPERACION if (pie and resaltar_pie) else None)


def entradas_vistas(manifiesto: dict) -> list[dict]:
    """El informe fotografico de una recepcion, vista por vista y emparejado.

    Cuando el acta trae las fotos de la salida, cada vista se imprime con la
    forma del formato —DESPACHO a la izquierda, RECEPCION a la derecha— en vez
    de como rejilla. Asi el informe fotografico y las OBSERVACIONES quedan
    juntos y con la misma forma, y ninguna foto se imprime dos veces.

    Sin ninguna foto de salida no hay nada que emparejar y el acta vuelve a la
    rejilla, que es la forma de la plantilla y ocupa la mitad de hojas.
    """
    if manifiesto["encabezado"].get("tipo_documento") != "RECEPCION":
        return []
    vistas = manifiesto.get("registro_fotografico", [])
    if not any(v.get("archivo_despacho") for v in vistas):
        return []
    entradas = []
    for foto in vistas:
        rotulo = foto.get("descripcion", "")
        entradas.append({
            "rotulo": rotulo,
            "foto_izq": foto.get("archivo_despacho"),
            "foto_der": foto.get("archivo"),
            # El rotulo va igual en las dos columnas: quien manda es el
            # encabezado DESPACHO / RECEPCION, como en el formato.
            "texto_izq": rotulo,
            "texto_der": rotulo,
            "pie": foto.get("observacion") or "",
        })
    return entradas


def entradas_danos(manifiesto: dict) -> list[dict]:
    """Solo los componentes observados o dañados: lo que sustenta un cobro."""
    if manifiesto["encabezado"].get("tipo_documento") != "RECEPCION":
        return []
    entradas = []
    for item in manifiesto.get("inspeccion_componentes", []):
        if item.get("estado") not in {"OBS", "D"}:
            continue
        nombre = item.get("item", "")
        etiqueta = textos.LEYENDA_ESTADO.get(item.get("estado"), "")
        entradas.append({
            "rotulo": nombre.upper(),
            "foto_izq": item.get("foto_despacho"),
            "foto_der": item.get("foto_recepcion"),
            "texto_izq": f"ANTES · {nombre.upper()}",
            "texto_der": f"DESPUÉS · {nombre.upper()} — {etiqueta.upper()}",
            "pie": item.get("observacion") or "",
        })
    return entradas


def _configurar_pagina(ws, ultima_fila: int) -> None:
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.24
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.print_area = f"A1:{layout.COL_ULTIMA}{ultima_fila}"
    ws.sheet_view.showGridLines = False


def construir_hoja_reporte(wb: Workbook, manifiesto: dict, raiz: Path,
                           avisos: list[str]):
    ws = wb.active
    ws.title = "REPORTE"
    fotos = manifiesto.get("registro_fotografico", [])
    consumibles = manifiesto.get("consumibles", [])

    vistas = entradas_vistas(manifiesto)
    danos = entradas_danos(manifiesto)

    tipo = manifiesto["encabezado"].get("tipo_documento", "RECEPCION")
    bandas = bandas_de_cierre(consumibles, tipo)

    # Emparejado o en rejilla, pero no las dos cosas: si las vistas se imprimen
    # emparejadas, la rejilla repetiria cada foto del retorno.
    n_bloques = 0 if vistas else max(1, (len(fotos) + 1) // 2)
    plan = layout.plan_secciones(n_bloques, len(consumibles), len(danos),
                                 bandas, len(vistas))
    fin = plan["ultima_fila"]

    _dimensionar(ws, fin, n_bloques, len(consumibles), bandas, len(vistas))
    _cabecera(ws, manifiesto["encabezado"], raiz, avisos)
    _seccion_pareada(ws, None, layout.FILA_VISTAS, vistas, raiz, avisos)
    _rejilla_fotografica(ws, fotos, raiz, avisos, pintar=not vistas)
    _bloques_consumibles(ws, consumibles, n_bloques, raiz, avisos, tipo, bandas,
                         len(vistas))
    _seccion_pareada(ws, layout.TITULO_DANOS, plan["fila_danos"],
                     danos, raiz, avisos, resaltar_pie=True)
    _saltos_de_pagina(ws, n_bloques, bandas, plan, len(danos), len(vistas))
    _configurar_pagina(ws, fin)
    return ws


# --------------------------------------------------------------------------- #
# Hojas auxiliares
# --------------------------------------------------------------------------- #
def _cabecera_auxiliar(ws, manifiesto: dict, titulo: str, n_columnas: int) -> int:
    """Encabezado de hoja auxiliar. Devuelve la fila donde empieza la tabla."""
    enc = manifiesto["encabezado"]
    ultima = get_column_letter(n_columnas)
    empresa = (enc.get("empresa") or "").strip()
    st.escribir(ws, f"A1:{ultima}1", f"{titulo} — {empresa}" if empresa else titulo,
                fuente=st.FUENTE_META, fill=st.FILL_CABECERA)
    detalle = (
        f"ACTA {enc.get('n_acta', '')}   |   {enc.get('tipo_documento', '')}   |   "
        f"{enc.get('fecha', '')}   |   EQUIPO {enc.get('codigo_equipo', '')} — "
        f"{enc.get('modelo_equipo', '')}   |   CLIENTE {enc.get('cliente', '').strip()}"
    )
    st.escribir(ws, f"A2:{ultima}2", detalle,
                fuente=st.FUENTE_VALOR, alineacion=st.IZQUIERDA_AJUSTADA)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 32
    return 4


def _hoja_tabular(wb, manifiesto: dict, titulo: str, encabezados: list[str],
                  filas: list[list], anchos: list[float]):
    ws = wb.create_sheet(titulo)
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    fila_encabezado = _cabecera_auxiliar(ws, manifiesto, titulo, len(encabezados))
    for i, texto in enumerate(encabezados, start=1):
        celda = ws.cell(row=fila_encabezado, column=i, value=texto)
        celda.font = st.FUENTE_ETIQUETA
        celda.fill = st.FILL_CABECERA
        celda.alignment = st.CENTRO_AJUSTADO
        celda.border = st.BORDE
    ws.row_dimensions[fila_encabezado].height = 28
    for r, fila in enumerate(filas, start=fila_encabezado + 1):
        for c, valor in enumerate(fila, start=1):
            celda = ws.cell(row=r, column=c, value=valor)
            celda.font = st.FUENTE_VALOR
            celda.alignment = st.IZQUIERDA_AJUSTADA if c > 1 else st.CENTRO
            celda.border = st.BORDE
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False
    return ws


def construir_hoja_inspeccion(wb: Workbook, manifiesto: dict):
    items = manifiesto.get("inspeccion_componentes", [])
    if not items:
        return None
    filas = [
        [i, it.get("item", ""), it.get("estado", ""),
         textos.LEYENDA_ESTADO.get(it.get("estado"), ""), it.get("observacion", "")]
        for i, it in enumerate(items, start=1)
    ]
    return _hoja_tabular(
        wb, manifiesto, "INSPECCIÓN",
        ["N°", "COMPONENTE INSPECCIONADO", "ESTADO", "SIGNIFICADO", "OBSERVACIÓN"],
        filas, [5, 42, 10, 18, 60],
    )


def construir_hoja_consumibles(wb: Workbook, manifiesto: dict):
    """Formato de consumibles: obligatorio en equipos moviles / autopropulsados."""
    enc = manifiesto["encabezado"]
    control = manifiesto.get("control_consumibles", [])
    if not control:
        control = textos.control_consumibles_vacio(enc.get("categoria", "generico"))

    filas = []
    for i, c in enumerate(control, start=1):
        despacho, recepcion = c.get("despacho"), c.get("recepcion")
        consumo = c.get("consumo")
        if (consumo is None and isinstance(despacho, (int, float))
                and isinstance(recepcion, (int, float))):
            consumo = round(despacho - recepcion, 2)
        filas.append([i, c.get("consumible", ""), c.get("unidad", ""),
                      despacho, recepcion, consumo,
                      c.get("estado", ""), c.get("observacion", "")])

    ws = _hoja_tabular(
        wb, manifiesto, "CONSUMIBLES",
        ["N°", "CONSUMIBLE / FLUIDO", "UNIDAD", "DESPACHO", "RECEPCIÓN",
         "CONSUMO", "ESTADO", "OBSERVACIÓN"],
        filas, [5, 34, 12, 12, 12, 12, 10, 46],
    )
    fila_firma = len(filas) + 7
    empresa = (manifiesto["encabezado"].get("empresa") or "").strip()
    rotulo_operador = f"FIRMA OPERADOR {empresa}".strip() if empresa else "FIRMA DEL OPERADOR"
    ws.cell(row=fila_firma, column=2, value=rotulo_operador).font = st.FUENTE_ETIQUETA
    ws.cell(row=fila_firma, column=5, value="FIRMA / V°B° CLIENTE").font = st.FUENTE_ETIQUETA
    for col in (2, 5):
        ws.cell(row=fila_firma + 3, column=col, value="_______________________________")
    return ws


def construir_hojas_guia(wb: Workbook, manifiesto: dict):
    for titulo, lineas in (
        ("GUÍA OPERADOR", textos.guia_operador(manifiesto)),
        ("GUÍA CLIENTE", textos.guia_cliente(manifiesto)),
    ):
        ws = wb.create_sheet(titulo)
        ws.column_dimensions["A"].width = 4
        ws.column_dimensions["B"].width = 110
        for i, (nivel, texto) in enumerate(lineas, start=1):
            celda = ws.cell(row=i, column=2, value=texto)
            celda.alignment = st.IZQUIERDA_AJUSTADA
            if nivel == "t":
                celda.font = st.FUENTE_META
                celda.fill = st.FILL_CABECERA
                ws.row_dimensions[i].height = 22
            elif nivel == "s":
                celda.font = st.FUENTE_ETIQUETA
            else:
                celda.font = st.FUENTE_VALOR
                ws.row_dimensions[i].height = max(15, 13 * (1 + len(texto) // 105))
        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_view.showGridLines = False


# --------------------------------------------------------------------------- #
# API
# --------------------------------------------------------------------------- #
def construir(manifiesto: dict, salida: str | Path, raiz: str | Path = ".",
              con_guias: bool = True) -> tuple[Path, list[str]]:
    """Escribe el libro Excel del reporte. Devuelve (ruta, avisos)."""
    raiz = Path(raiz)
    salida = Path(salida)
    avisos: list[str] = []

    wb = Workbook()
    construir_hoja_reporte(wb, manifiesto, raiz, avisos)
    construir_hoja_inspeccion(wb, manifiesto)
    construir_hoja_consumibles(wb, manifiesto)
    if con_guias:
        construir_hojas_guia(wb, manifiesto)

    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)
    return salida, avisos
